"""
Warframe Market Sell Sync - Configuration & Design System
=========================================================
Defines design tokens, fonts, fills, alignments, borders, column widths,
and global application constants.
"""

import sys
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

# Configure safe stdout encoding on Windows
if sys.platform == "win32" and hasattr(sys.stdout, "reconfigure"):
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except Exception:
        pass

# File & API Constants
CACHE_FILE = "items_cache.json"
DEFAULT_EXCEL = "Warframe Sell Stats.xlsm"
SHEET_NAME = "Sheet1"
API_BASE_URL = "https://api.warframe.market/v2"
API_CALL_DELAY_SEC = 0.6
DEFAULT_EXPORT_DIR_PLACEHOLDER = "<PATH_TO_WARFRAME_FOLDER>"

BROWSER_HEADERS = {
    "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/148.0.0.0 Safari/537.36 Edg/148.0.0.0",
    "Accept": "application/json",
    "Content-Type": "application/json",
    "Platform": "pc",
    "Language": "en",
    "Crossplay": "true"
}

# Typography & Color Palette - Modern Sleek Dark Mode
FONT_NAME = "Segoe UI"
FONT_HEADER = Font(name=FONT_NAME, size=11, bold=True, color="FFFFFF")
FONT_SUBHEADER = Font(name=FONT_NAME, size=10, bold=True, color="E2E8F0")
FONT_DATA = Font(name=FONT_NAME, size=10, bold=False, color="F8FAFC")
FONT_ITEM_NAME = Font(name=FONT_NAME, size=10, bold=True, color="FFFFFF")
FONT_PRICE = Font(name=FONT_NAME, size=10, bold=True, color="FBBF24")        # Platinum Gold
FONT_STOCK = Font(name=FONT_NAME, size=10, bold=True, color="F8FAFC")        # Crisp White
FONT_CHECKED = Font(name=FONT_NAME, size=11, bold=True, color="34D399")      # Emerald Green (Visible)
FONT_UNCHECKED = Font(name=FONT_NAME, size=11, bold=True, color="64748B")    # Dim Slate (Hidden)
FONT_QTY = Font(name=FONT_NAME, size=10, bold=False, color="93C5FD")          # Ice Blue for quantities
FONT_REV = Font(name=FONT_NAME, size=10, bold=True, color="34D399")           # Emerald Green for revenue
FONT_TOTAL = Font(name=FONT_NAME, size=11, bold=True, color="38BDF8")         # Cyan Accent for Totals
FONT_TOTAL_LABEL = Font(name=FONT_NAME, size=11, bold=True, color="FFFFFF")
FONT_MONO = Font(name="Consolas", size=9, color="FDE68A")                     # Light Gold for JWT Token

# Dark Theme Fills
FILL_CANVAS = PatternFill(start_color="0B0F19", end_color="0B0F19", fill_type="solid")      # Deepest Midnight
FILL_ROW_ODD = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")     # Sleek Slate 800
FILL_ROW_EVEN = PatternFill(start_color="162032", end_color="162032", fill_type="solid")    # Darker Slate 850
FILL_ZEBRA_EVEN = FILL_ROW_EVEN
FILL_HEADER_DARK = PatternFill(start_color="0B0F19", end_color="0B0F19", fill_type="solid") # Obsidian Slate
FILL_HEADER_BLUE = PatternFill(start_color="0369A1", end_color="0369A1", fill_type="solid") # Ocean Blue
FILL_HEADER_GREEN = PatternFill(start_color="065F46", end_color="065F46", fill_type="solid")# Deep Emerald
FILL_SUBHEADER = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")   # Slate 800
FILL_TOTAL = PatternFill(start_color="0B0F19", end_color="0B0F19", fill_type="solid")       # Obsidian Totals
FILL_TOKEN = PatternFill(start_color="241B13", end_color="241B13", fill_type="solid")       # Bronze Token Box
FILL_CARD_DARK = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")   # Sidebar Card

# Cell Alignments
ALIGN_LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
ALIGN_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
ALIGN_RIGHT = Alignment(horizontal="right", vertical="center", wrap_text=True)

# Borders
BORDER_THIN = Side(style="thin", color="334155")      # Slate 700 Border
BORDER_MEDIUM = Side(style="medium", color="475569")  # Slate 600
BORDER_DOUBLE = Side(style="double", color="38BDF8")  # Cyan Accent Double Line
BORDER_AMBER = Side(style="thin", color="D97706")     # Amber Accent for Token Box

BORDER_CELL = Border(left=BORDER_THIN, right=BORDER_THIN, top=BORDER_THIN, bottom=BORDER_THIN)
BORDER_TOTAL = Border(top=BORDER_MEDIUM, bottom=BORDER_DOUBLE, left=BORDER_THIN, right=BORDER_THIN)
BORDER_TOKEN_BOX = Border(left=BORDER_AMBER, right=BORDER_AMBER, top=BORDER_AMBER, bottom=BORDER_AMBER)

# Standard Column Widths
COLUMN_WIDTHS = {
    "A": 34, # Item/Mod Name
    "B": 12, # Price
    "C": 12, # Stock
    "D": 10, # Visible Checkbox
    "E": 16, # Quantity Sold: Current Session
    "F": 14, # Quantity Sold: All Time
    "G": 18, # Revenue Generated: Current Session
    "H": 16, # Revenue Generated: All Time
    "I": 28, # Actions Gutter
    "J": 80  # Sidebar Controls
}
