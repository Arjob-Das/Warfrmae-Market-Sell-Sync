"""
Warframe Market Sell Sync - Market Updater & Diff Engine
========================================================
Reads current pricing, evaluated stock, and visibility from Excel,
compares against active market orders, and pushes updates to Warframe.market.
"""

import os
import time
import requests
import openpyxl
from typing import List, Tuple

from .config import (
    API_BASE_URL, API_CALL_DELAY_SEC, SHEET_NAME
)
from .sheet_layout import get_open_excel_workbook, prompt_jwt_token
from .api import get_auth_headers_and_cookies


def read_excel_prices_and_stock(excel_file: str) -> List[dict]:
    """Reads items, prices, evaluated stock, and visibility checkbox from Excel sheet."""
    if not os.path.exists(excel_file):
        return []

    app, wb_com = get_open_excel_workbook(excel_file)
    if wb_com:
        try:
            wb_com.Save()
        except Exception:
            pass

    wb_form = openpyxl.load_workbook(excel_file, data_only=False)
    ws_form = wb_form[SHEET_NAME] if SHEET_NAME in wb_form.sheetnames else wb_form.active
    wb_val = openpyxl.load_workbook(excel_file, data_only=True)
    ws_val = wb_val[SHEET_NAME] if SHEET_NAME in wb_val.sheetnames else wb_val.active

    items = []
    for r in range(3, ws_form.max_row + 1):
        name_val = ws_form.cell(row=r, column=1).value
        price_val = ws_val.cell(row=r, column=2).value
        if not name_val or str(name_val).strip().lower() == "total":
            continue
        try:
            p = int(price_val)
        except (ValueError, TypeError):
            continue

        c_formula = str(ws_form.cell(row=r, column=3).value or "").strip()
        e_val = ws_val.cell(row=r, column=5).value or 0
        try:
            sold_val = int(e_val)
        except (ValueError, TypeError):
            sold_val = 0

        stock_val = None
        base_num = 1
        if c_formula.startswith("="):
            parts = c_formula[1:].split("-")
            if parts:
                try:
                    base_num = int(parts[0].strip())
                    stock_val = max(0, base_num - sold_val)
                except ValueError:
                    pass
        if stock_val is None:
            c_eval = ws_val.cell(row=r, column=3).value
            try:
                stock_val = max(0, int(float(c_eval)))
            except (ValueError, TypeError):
                stock_val = 1
            base_num = stock_val + sold_val

        d_val = ws_val.cell(row=r, column=4).value
        if d_val is None or str(d_val).strip() == "":
            is_vis = True
        else:
            d_str = str(d_val).strip()
            is_vis = d_str in ("☑", "True", "true", "TRUE", "1", "yes", "Yes")

        # When effective stock is 0 or less, the item is depleted / out of stock,
        # so effective visibility MUST be False (delisted).
        if stock_val <= 0:
            is_vis = False

        items.append({
            "name": str(name_val).strip(),
            "price": p,
            "stock": stock_val,
            "visible": is_vis,
            "base_stock": base_num,
            "curr_sold": sold_val,
            "row": r
        })
    wb_form.close()
    wb_val.close()
    return items


def read_excel_prices(excel_file: str) -> List[dict]:
    """Compatibility wrapper calling read_excel_prices_and_stock."""
    return read_excel_prices_and_stock(excel_file)


def diff_prices_and_stock(excel_items: List[dict], market_orders: List[dict]) -> Tuple[List[dict], List[dict], List[dict]]:
    """Compares Excel prices, stock, and visibility with active market orders for two-way sync."""
    market_lookup = {}
    for o in market_orders:
        market_lookup[o["name"].strip().lower()] = o
        base_k = o["base_name"].strip().lower()
        if base_k not in market_lookup:
            market_lookup[base_k] = o

    changed = []
    unchanged = []
    unmatched = []

    for ex in excel_items:
        ex_name = ex["name"].strip()
        ex_key = ex_name.lower()
        ex_price = ex["price"]
        ex_stock = ex["stock"]
        # If stock is 0 or less, visibility must be False
        ex_visible = False if ex_stock <= 0 else ex["visible"]

        order = market_lookup.get(ex_key)
        if not order:
            base_name = ex_name.split(" (Rank")[0].strip().lower()
            order = market_lookup.get(base_name)

        if not order:
            unmatched.append(ex)
            continue

        price_diff = (order["price"] != ex_price)

        # Warframe.market order quantity cannot be 0 (API requires quantity >= 1).
        # If ex_stock is 0 or less, or if ex_visible is False:
        # We maintain order quantity at 1, but set visible to False (delisting it from market).
        # If ex_stock > 0 and ex_visible is True:
        # Target quantity is ex_stock, and target visibility is True.
        if ex_stock <= 0 or not ex_visible:
            target_qty = 1
            target_vis = False
        else:
            target_qty = ex_stock
            target_vis = True

        stock_diff = (order["quantity"] != target_qty and target_vis)
        visible_diff = (bool(order.get("visible", True)) != target_vis)

        if price_diff or stock_diff or visible_diff:
            changed.append({
                "name": ex_name,
                "order_id": order["order_id"],
                "old_price": order["price"],
                "new_price": ex_price,
                "old_quantity": order["quantity"],
                "new_quantity": target_qty,
                "old_visible": bool(order.get("visible", True)),
                "new_visible": target_vis,
                "price_diff": price_diff,
                "stock_diff": stock_diff,
                "visible_diff": visible_diff,
                "rank": order["rank"],
                "visible": target_vis,
                "excel_stock": ex_stock,
                "excel_row": ex.get("row")
            })
        else:
            unchanged.append({
                "name": ex_name,
                "price": ex_price,
                "stock": ex_stock,
                "visible": target_vis
            })

    return changed, unchanged, unmatched


def diff_prices(excel_items: List[dict], market_orders: List[dict]) -> Tuple[List[dict], List[dict], List[dict]]:
    """Compatibility wrapper calling diff_prices_and_stock."""
    return diff_prices_and_stock(excel_items, market_orders)


def display_price_diff(changed: List[dict], unchanged: List[dict], unmatched: List[dict]) -> None:
    """Prints a formatted diff table of price, stock, and visibility changes."""
    print("\n" + "=" * 86)
    print("  WARFRAME MARKET PRICE, STOCK & VISIBILITY UPDATE PREVIEW")
    print("=" * 86)

    if changed:
        print(f"\n  {len(changed)} CHANGE(S) DETECTED:")
        print("  " + "-" * 82)
        print(f"  {'Item Name':<30} {'Price (Mkt -> XL)':<20} {'Stock (Mkt -> XL)':<18} {'Visible'}")
        print("  " + "-" * 82)
        for it in changed:
            if it["price_diff"]:
                diff = it["new_price"] - it["old_price"]
                sign = "+" if diff > 0 else "-"
                price_str = f"{it['old_price']} -> {it['new_price']} ({sign}{abs(diff)}p)"
            else:
                price_str = f"{it['new_price']} (same)"

            if it["stock_diff"]:
                sdiff = it["new_quantity"] - it["old_quantity"]
                ssign = "+" if sdiff > 0 else "-"
                stock_str = f"{it['old_quantity']} -> {it['new_quantity']} ({ssign}{abs(sdiff)})"
            elif it.get("excel_stock", 1) <= 0:
                stock_str = f"{it['old_quantity']} -> 0 (Depleted)"
            else:
                stock_str = f"{it['new_quantity']} (same)"

            if it["visible_diff"]:
                old_v = "☑ Vis" if it["old_visible"] else "☐ Hid"
                new_v = "☑ Vis" if it["new_visible"] else "☐ Hid"
                vis_str = f"{old_v} -> {new_v}"
            else:
                vis_str = "☑ Vis" if it["new_visible"] else "☐ Hid"

            print(f"  {it['name']:<30} {price_str:<20} {stock_str:<18} {vis_str}")
        print("  " + "-" * 82)
    else:
        print("\n  No price, stock, or visibility changes detected. Excel matches warframe.market!")

    if unchanged:
        print(f"  {len(unchanged)} item(s) already match market price, stock, and visibility.")
    if unmatched:
        print(f"  {len(unmatched)} item(s) in Excel were not found in market orders.")
    print("=" * 86)


def push_price_updates(
    jwt_token: str,
    changes: List[dict],
    excel_file: str
) -> Tuple[int, int]:
    """Pushes price, stock, and visibility updates via PATCH /v2/order/{order_id}."""
    headers, cookies = get_auth_headers_and_cookies(jwt_token)
    success = 0
    failed = 0
    succeeded_changes = []

    for i, ch in enumerate(changes):
        order_id = ch["order_id"]
        new_price = ch["new_price"]
        new_qty = ch["new_quantity"]
        item_name = ch["name"]

        payload = {
            "platinum": new_price,
            "quantity": new_qty,
            "visible": ch["visible"]
        }
        if ch["rank"] is not None:
            payload["rank"] = ch["rank"]

        url = f"{API_BASE_URL}/order/{order_id}"
        vis_tag = "visible" if ch["visible"] else "HIDDEN"
        try:
            resp = requests.patch(url, headers=headers, cookies=cookies, json=payload, timeout=10)
            if resp.status_code in (200, 201, 204):
                print(f"  [{i+1}/{len(changes)}] Updated '{item_name}': {new_price} plat, {new_qty} stock, {vis_tag}")
                success += 1
                succeeded_changes.append(ch)
            elif resp.status_code == 401:
                print(f"  [!] JWT expired while updating '{item_name}'.")
                new_tok = prompt_jwt_token(excel_file)
                if new_tok:
                    jwt_token = new_tok
                    headers, cookies = get_auth_headers_and_cookies(jwt_token)
                    retry = requests.patch(url, headers=headers, cookies=cookies, json=payload, timeout=10)
                    if retry.status_code in (200, 201, 204):
                        print(f"  [{i+1}/{len(changes)}] Updated '{item_name}': {new_price} plat, {new_qty} stock, {vis_tag}")
                        success += 1
                        succeeded_changes.append(ch)
                        continue
                failed += 1
                break
            else:
                print(f"  [{i+1}/{len(changes)}] FAILED '{item_name}': HTTP {resp.status_code} - {resp.text[:100]}")
                failed += 1
        except Exception as e:
            print(f"  [{i+1}/{len(changes)}] FAILED '{item_name}': {e}")
            failed += 1

        if i < len(changes) - 1:
            time.sleep(API_CALL_DELAY_SEC)

    # Synchronize Excel Column D in workbook for pushed items
    if succeeded_changes and os.path.exists(excel_file):
        app, wb_com = get_open_excel_workbook(excel_file)
        if wb_com:
            try:
                ws = wb_com.Sheets(1)
                for ch in succeeded_changes:
                    r = ch.get("excel_row")
                    if r and ch.get("new_visible") is not None:
                        if ch["new_visible"]:
                            ws.Cells(r, 4).Value = chr(0x2611)  # ☑
                            ws.Cells(r, 4).Font.Color = 52 + (211 * 256) + (153 * 65536)  # RGB(52, 211, 153)
                        else:
                            ws.Cells(r, 4).Value = chr(0x2610)  # ☐
                            ws.Cells(r, 4).Font.Color = 100 + (116 * 256) + (139 * 65536)  # RGB(100, 116, 139)
                wb_com.Save()
            except Exception:
                pass
        else:
            try:
                from .config import FONT_CHECKED, FONT_UNCHECKED
                wb = openpyxl.load_workbook(excel_file, keep_vba=True)
                ws = wb[SHEET_NAME] if SHEET_NAME in wb.sheetnames else wb.active
                for ch in succeeded_changes:
                    r = ch.get("excel_row")
                    if r and ch.get("new_visible") is not None:
                        if ch["new_visible"]:
                            ws.cell(row=r, column=4, value="☑")
                            ws.cell(row=r, column=4).font = FONT_CHECKED
                        else:
                            ws.cell(row=r, column=4, value="☐")
                            ws.cell(row=r, column=4).font = FONT_UNCHECKED
                wb.save(excel_file)
                wb.close()
            except Exception:
                pass

    return success, failed
