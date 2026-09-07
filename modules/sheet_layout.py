"""
Warframe Market Sell Sync - Sheet Layout & Formatting Engine
============================================================
Handles worksheet structure, 2-row merged headers, row formulas,
dark-mode styling, sidebar credentials, and migration of legacy layouts.
"""

import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border
from typing import Optional, Tuple

from .config import (
    FONT_NAME, FONT_HEADER, FONT_SUBHEADER, FONT_DATA, FONT_ITEM_NAME,
    FONT_PRICE, FONT_STOCK, FONT_CHECKED, FONT_UNCHECKED, FONT_QTY,
    FONT_REV, FONT_TOTAL, FONT_TOTAL_LABEL, FONT_MONO,
    FILL_CANVAS, FILL_ROW_ODD, FILL_ROW_EVEN, FILL_HEADER_DARK,
    FILL_HEADER_BLUE, FILL_HEADER_GREEN, FILL_SUBHEADER, FILL_TOTAL,
    FILL_TOKEN, FILL_CARD_DARK, ALIGN_LEFT, ALIGN_CENTER, ALIGN_RIGHT,
    BORDER_CELL, BORDER_TOTAL, BORDER_TOKEN_BOX, COLUMN_WIDTHS,
    SHEET_NAME, DEFAULT_EXPORT_DIR_PLACEHOLDER
)


def get_open_excel_workbook(excel_file: str):
    """Returns (excel_app, workbook) if excel_file is currently open in Excel COM, else (None, None)."""
    try:
        import win32com.client
        excel = win32com.client.GetActiveObject("Excel.Application")
        target = os.path.abspath(excel_file).lower()
        for wb in excel.Workbooks:
            if os.path.abspath(wb.FullName).lower() == target:
                return excel, wb
    except Exception:
        pass
    return None, None


def get_credentials_from_excel(excel_file: str) -> Tuple[str, str]:
    """Reads Username from cell J2 (fallback I2, H2) and JWT Token from cell J4 (fallback I4, H4)."""
    username = ""
    jwt_token = ""
    if not os.path.exists(excel_file):
        return username, jwt_token

    # Check active Excel first
    app, wb_com = get_open_excel_workbook(excel_file)
    if wb_com:
        try:
            ws = wb_com.Sheets(1)
            h2_val = str(ws.Range("J2").Value or ws.Range("I2").Value or ws.Range("H2").Value or "").strip()
            if h2_val and h2_val.lower() not in ("none", "enter your username here"):
                username = h2_val
            h4_val = str(ws.Range("J4").Value or ws.Range("I4").Value or ws.Range("H4").Value or "").strip()
            if h4_val and h4_val.lower() not in ("none", "paste your jwt token here"):
                if h4_val.lower().startswith("bearer "):
                    h4_val = h4_val[7:].strip()
                elif h4_val.lower().startswith("jwt "):
                    h4_val = h4_val[4:].strip()
                if len(h4_val) > 20:
                    jwt_token = h4_val
            return username, jwt_token
        except Exception:
            pass

    try:
        wb = openpyxl.load_workbook(excel_file, data_only=True, keep_vba=True)
        ws = wb[SHEET_NAME] if SHEET_NAME in wb.sheetnames else wb.active

        h2_val = str(ws["J2"].value or ws["I2"].value or ws["H2"].value or "").strip()
        if h2_val and h2_val.lower() not in ("none", "enter your username here"):
            username = h2_val

        h4_val = str(ws["J4"].value or ws["I4"].value or ws["H4"].value or "").strip()
        if h4_val and h4_val.lower() not in ("none", "paste your jwt token here"):
            if h4_val.lower().startswith("bearer "):
                h4_val = h4_val[7:].strip()
            elif h4_val.lower().startswith("jwt "):
                h4_val = h4_val[4:].strip()
            if len(h4_val) > 20:
                jwt_token = h4_val

        wb.close()
    except Exception:
        pass
    return username, jwt_token


def save_credentials_to_excel(excel_file: str, username: str = "", jwt_token: str = "") -> None:
    """Writes Username to cell J2 and JWT Token to cell J4 in the Excel workbook."""
    if not os.path.exists(excel_file):
        return

    # Check active Excel first to prevent file lock PermissionError
    app, wb_com = get_open_excel_workbook(excel_file)
    if wb_com:
        try:
            ws = wb_com.Sheets(1)
            if username:
                ws.Range("J2").Value = username
            if jwt_token:
                clean_tok = jwt_token.strip()
                if clean_tok.lower().startswith("bearer "):
                    clean_tok = clean_tok[7:].strip()
                elif clean_tok.lower().startswith("jwt "):
                    clean_tok = clean_tok[4:].strip()
                ws.Range("J4").Value = clean_tok
            wb_com.Save()
            return
        except Exception:
            pass

    try:
        wb = openpyxl.load_workbook(excel_file, keep_vba=True)
        ws = wb[SHEET_NAME] if SHEET_NAME in wb.sheetnames else wb.active

        if username:
            ws["J2"] = username
        if jwt_token:
            clean_tok = jwt_token.strip()
            if clean_tok.lower().startswith("bearer "):
                clean_tok = clean_tok[7:].strip()
            elif clean_tok.lower().startswith("jwt "):
                clean_tok = clean_tok[4:].strip()
            ws["J4"] = clean_tok

        wb.save(excel_file)
        wb.close()
    except Exception as e:
        print(f"[!] Warning: Could not update credentials in Excel: {e}")


def prompt_jwt_token(excel_file: str) -> str:
    """Prompts user for a fresh JWT token and writes it directly into Excel cell J4."""
    print("\n" + "=" * 70)
    print("  JWT TOKEN REQUIRED / EXPIRED - Warframe.market Authentication")
    print("=" * 70)
    print("How to get a fresh token from your browser (Chrome / Edge):")
    print("  1. Open https://warframe.market (ensure you are logged in)")
    print("  2. Press F12 -> Go to 'Application' or 'Storage' tab")
    print("  3. Under Storage -> Cookies -> https://warframe.market")
    print("  4. Find the cookie named 'JWT'")
    print("  5. Copy the value (starts with 'eyJ...')")
    print("  Note: You can also paste this directly into cell J4 in Excel!")
    print("=" * 70)

    token = input("\nPaste your JWT token here: ").strip()
    if token.lower().startswith("bearer "):
        token = token[7:].strip()
    elif token.lower().startswith("jwt "):
        token = token[4:].strip()

    if token:
        save_credentials_to_excel(excel_file, jwt_token=token)
        print(f"[+] Token saved to cell J4 of '{excel_file}'.")
    return token


def get_export_dir_from_excel(excel_file: str) -> Optional[str]:
    """Reads Export Directory from cell J16 (fallback I16, H16)."""
    if not os.path.exists(excel_file):
        return None

    app, wb_com = get_open_excel_workbook(excel_file)
    if wb_com:
        try:
            ws = wb_com.Sheets(1)
            h16_val = str(ws.Range("J16").Value or ws.Range("I16").Value or ws.Range("H16").Value or "").strip()
            if h16_val and not h16_val.startswith("<") and h16_val.lower() not in ("none", ""):
                return h16_val
        except Exception:
            pass

    try:
        wb = openpyxl.load_workbook(excel_file, data_only=True, keep_vba=True)
        ws = wb[SHEET_NAME] if SHEET_NAME in wb.sheetnames else wb.active
        h16_val = str(ws["J16"].value or ws["I16"].value or ws["H16"].value or "").strip()
        wb.close()
        if h16_val and not h16_val.startswith("<") and h16_val.lower() not in ("none", ""):
            return h16_val
    except Exception:
        pass
    return None


def save_export_dir_to_excel(excel_file: str, export_dir: str) -> None:
    """Writes Export Directory path to cell J16 in the Excel workbook."""
    if not os.path.exists(excel_file) or not export_dir:
        return

    app, wb_com = get_open_excel_workbook(excel_file)
    if wb_com:
        try:
            ws = wb_com.Sheets(1)
            ws.Range("J16").Value = export_dir
            wb_com.Save()
            return
        except Exception:
            pass

    try:
        wb = openpyxl.load_workbook(excel_file, keep_vba=True)
        ws = wb[SHEET_NAME] if SHEET_NAME in wb.sheetnames else wb.active
        ws["J16"] = export_dir
        wb.save(excel_file)
        wb.close()
    except Exception as e:
        print(f"[!] Warning: Could not update export directory in Excel: {e}")


def safe_merge(ws, cell_range: str) -> None:
    """Safely merges cells if range is not already merged."""
    existing = [str(r) for r in ws.merged_cells.ranges]
    if cell_range not in existing:
        ws.merge_cells(cell_range)


def migrate_sheet_layout_if_needed(ws) -> bool:
    """
    Detects if the sheet is in older formats and migrates to 8-column layout:
    A: Item Name, B: Price, C: Stock, D: Visible, E: Qty Session, F: Qty All, G: Rev Session, H: Rev All, I: Gutter, J: Sidebar
    """
    c1 = ws["C1"].value
    d1 = ws["D1"].value
    c1_str = str(c1 or "").strip().lower()
    d1_str = str(d1 or "").strip().lower()

    # Case 1: 6-column sheet (C1 is Quantity Sold)
    if "quantity" in c1_str:
        ws.insert_cols(3, amount=2)
        for r in range(3, ws.max_row + 1):
            if ws.cell(row=r, column=1).value and str(ws.cell(row=r, column=1).value).lower() != "total":
                ws.cell(row=r, column=3, value=1)
                ws.cell(row=r, column=4, value="☑")
        return True

    # Case 2: 7-column sheet (C1 is Stock, D1 is Quantity Sold)
    if "stock" in c1_str and "quantity" in d1_str:
        ws.insert_cols(4, amount=1)
        for r in range(3, ws.max_row + 1):
            if ws.cell(row=r, column=1).value and str(ws.cell(row=r, column=1).value).lower() != "total":
                ws.cell(row=r, column=4, value="☑")
        return True

    return False


def initialize_sheet_structure(ws) -> None:
    """Builds standard 2-row table headers, merged groups, labels, Stock, and Visible column."""
    for rng in list(ws.merged_cells.ranges):
        if rng.min_row <= 2:
            try:
                ws.unmerge_cells(str(rng))
            except Exception:
                pass

    ws["A1"] = "Item/Mod Name"
    ws["B1"] = "Price"
    ws["C1"] = "Stock"
    ws["D1"] = "Visible"
    ws["E1"] = "Quantity Sold"
    ws["F1"] = None
    ws["G1"] = "Revenue Generated"
    ws["H1"] = None

    ws["A2"] = None
    ws["B2"] = None
    ws["C2"] = None
    ws["D2"] = None
    ws["E2"] = "Current Session"
    ws["F2"] = "All Time"
    ws["G2"] = "Current Session"
    ws["H2"] = "All Time"

    safe_merge(ws, "A1:A2")
    safe_merge(ws, "B1:B2")
    safe_merge(ws, "C1:C2")
    safe_merge(ws, "D1:D2")
    safe_merge(ws, "E1:F1")
    safe_merge(ws, "G1:H1")

    # Header cell styling
    for col in range(1, 9):
        cell1 = ws.cell(row=1, column=col)
        cell1.alignment = ALIGN_CENTER
        cell1.border = BORDER_CELL
        if col in (1, 2, 3, 4):
            cell1.fill = FILL_HEADER_DARK
            cell1.font = FONT_HEADER
        elif col in (5, 6):
            cell1.fill = FILL_HEADER_BLUE
            cell1.font = FONT_HEADER
        elif col in (7, 8):
            cell1.fill = FILL_HEADER_GREEN
            cell1.font = FONT_HEADER

        cell2 = ws.cell(row=2, column=col)
        cell2.alignment = ALIGN_CENTER
        cell2.border = BORDER_CELL
        cell2.font = FONT_SUBHEADER
        if col in (1, 2, 3, 4):
            cell2.fill = FILL_HEADER_DARK
        else:
            cell2.fill = FILL_SUBHEADER

    # I2: Linked Actions Header (Gutter)
    ws["I2"] = "Actions / Update All Time Revenue:"
    ws["I2"].font = Font(name=FONT_NAME, size=10, bold=True, color="38BDF8", underline="single")
    ws["I2"].fill = FILL_CARD_DARK
    ws["I2"].alignment = ALIGN_CENTER
    ws["I2"].border = BORDER_CELL

    ws.row_dimensions[1].height = 26
    ws.row_dimensions[2].height = 22

    try:
        if ws.views.sheetView:
            ws.views.sheetView[0].showGridLines = True
    except Exception:
        pass


def apply_row_formulas_and_styling(
    ws,
    row_idx: int,
    base_stock: Optional[int] = None,
    is_visible: Optional[bool] = None,
    base_all_time_rev: Optional[int] = None,
    is_zebra: bool = False
) -> None:
    """
    Applies dynamic Excel formulas:
      - Stock (Col C): ={base_stock}-E{row_idx}
      - Visible (Col D): '☑' or '☐'
      - Current Session Rev (Col G): =B{row_idx}*E{row_idx}
      - All Time Rev (Col H): ={base_all_time_rev}+G{row_idx} (Preserves historical revenue!)
    """
    # 1. Determine base stock
    if base_stock is None:
        c_val = ws.cell(row=row_idx, column=3).value
        parsed_base = 1
        if c_val is not None:
            c_str = str(c_val).strip()
            if c_str.startswith("="):
                parts = c_str[1:].split("-")
                if parts:
                    try:
                        parsed_base = int(parts[0].strip())
                    except ValueError:
                        parsed_base = 1
            else:
                try:
                    parsed_base = int(float(c_str))
                except ValueError:
                    parsed_base = 1
        base_stock = max(0, parsed_base)

    # 2. Determine visibility
    if is_visible is None:
        d_val = ws.cell(row=row_idx, column=4).value
        if d_val is None or str(d_val).strip() == "":
            is_visible = (base_stock > 0)
        else:
            d_str = str(d_val).strip()
            is_visible = d_str in ("☑", "True", "true", "TRUE", "1", "yes", "Yes")

    # If effective stock is 0 or less, visibility must be unchecked
    try:
        e_sold = int(ws.cell(row=row_idx, column=5).value or 0)
    except (ValueError, TypeError):
        e_sold = 0
    if (base_stock - e_sold) <= 0:
        is_visible = False

    # 3. Determine base all-time revenue
    if base_all_time_rev is None:
        h_cell = ws.cell(row=row_idx, column=8)
        h_val = str(h_cell.value or "").strip()
        parsed_base_rev = 0
        if h_val.startswith("="):
            parts = h_val[1:].split("+")
            if parts and len(parts) >= 2:
                try:
                    parsed_base_rev = int(parts[0].strip())
                except ValueError:
                    parsed_base_rev = 0
            else:
                # If old formula was =B*F, try to compute current all-time sold * price
                try:
                    f_val = int(ws.cell(row=row_idx, column=6).value or 0)
                    p_val = int(ws.cell(row=row_idx, column=2).value or 0)
                    parsed_base_rev = f_val * p_val
                except Exception:
                    parsed_base_rev = 0
        else:
            try:
                parsed_base_rev = int(float(h_val))
            except ValueError:
                parsed_base_rev = 0
        base_all_time_rev = max(0, parsed_base_rev)

    # Set cell formulas
    ws[f"C{row_idx}"] = f"={base_stock}-E{row_idx}"
    ws[f"D{row_idx}"] = "☑" if is_visible else "☐"
    ws[f"G{row_idx}"] = f"=B{row_idx}*E{row_idx}"
    ws[f"H{row_idx}"] = f"={base_all_time_rev}+G{row_idx}"

    row_fill = FILL_ROW_EVEN if is_zebra else FILL_ROW_ODD
    ws.row_dimensions[row_idx].height = 22
    for col_idx in range(1, 9):
        cell = ws.cell(row=row_idx, column=col_idx)
        cell.border = BORDER_CELL
        cell.fill = row_fill

        if col_idx == 1:
            cell.font = FONT_ITEM_NAME
            cell.alignment = ALIGN_LEFT
        elif col_idx == 2:
            cell.font = FONT_PRICE
            cell.alignment = ALIGN_RIGHT
            cell.number_format = "#,##0"
        elif col_idx == 3:
            cell.font = FONT_STOCK
            cell.alignment = ALIGN_RIGHT
            cell.number_format = "#,##0"
        elif col_idx == 4:
            cell.font = FONT_CHECKED if is_visible else FONT_UNCHECKED
            cell.alignment = ALIGN_CENTER
        elif col_idx in (5, 6):
            cell.font = FONT_QTY
            cell.alignment = ALIGN_RIGHT
            cell.number_format = "#,##0"
        elif col_idx in (7, 8):
            cell.font = FONT_REV
            cell.alignment = ALIGN_RIGHT
            cell.number_format = "#,##0"


def populate_sidebar_column(ws, username: str = "", jwt_token: str = "", export_dir: str = "") -> None:
    """Sets up Column J (with Column I gutter) for credentials, actions, and export settings."""
    max_r = max(ws.max_row or 0, 30)
    for r in range(1, max_r + 5):
        if r != 2:
            cell_i = ws.cell(row=r, column=9)
            if not isinstance(cell_i, openpyxl.cell.cell.MergedCell):
                cell_i.value = None

    # J1: Username Header
    ws["J1"] = "Warframe.market Account:"
    ws["J1"].font = Font(name=FONT_NAME, size=10, bold=True, color="FFFFFF")
    ws["J1"].fill = FILL_HEADER_DARK
    ws["J1"].alignment = ALIGN_LEFT

    # J2: Username Cell
    existing_j2 = str(ws["J2"].value or ws["I2"].value or ws["H2"].value or "").strip()
    u_val = username if username else (existing_j2 if existing_j2 and existing_j2.lower() != "none" else "")
    ws["J2"] = u_val
    ws["J2"].font = Font(name=FONT_NAME, size=10, bold=True, color="F8FAFC")
    ws["J2"].alignment = ALIGN_LEFT
    ws["J2"].fill = FILL_CARD_DARK
    ws["J2"].border = BORDER_CELL

    # J3: JWT Header
    ws["J3"] = "Stored JWT Token (warframe.market):"
    ws["J3"].font = Font(name=FONT_NAME, size=10, bold=True, color="FFFFFF")
    ws["J3"].fill = FILL_HEADER_DARK
    ws["J3"].alignment = ALIGN_LEFT

    # J4: JWT Cell
    existing_j4 = str(ws["J4"].value or ws["I4"].value or ws["H4"].value or "").strip()
    tok_val = jwt_token if jwt_token else (existing_j4 if existing_j4 and len(existing_j4) > 20 else "Paste your JWT token here")
    ws["J4"] = tok_val
    ws["J4"].font = FONT_MONO
    ws["J4"].fill = FILL_TOKEN
    ws["J4"].alignment = Alignment(horizontal="left", vertical="center", wrap_text=False)
    ws["J4"].border = BORDER_TOKEN_BOX

    # J5: Note
    ws["J5"] = "^ Stored credentials. Read directly by Python. Paste fresh JWT above if expired."
    ws["J5"].font = Font(name=FONT_NAME, size=8, italic=True, color="94A3B8")
    ws["J5"].fill = FILL_CANVAS

    # J6: Actions Header
    ws["J6"] = "One-Click Macro Actions (Click Cell or Button to Run):"
    ws["J6"].font = Font(name=FONT_NAME, size=10, bold=True, color="FFFFFF")
    ws["J6"].fill = FILL_HEADER_BLUE
    ws["J6"].alignment = ALIGN_LEFT

    # Action Items (Rows 7, 9, 11, 13)
    actions = [
        (7, "▶ 1. Sync from Market", "38BDF8", "0C2D48"),
        (9, "⬆ 2. Push Prices & Stock", "34D399", "064E3B"),
        (11, "🔄 3. Update All Time Revenue", "FBBF24", "3D1A04"),
        (13, "⚡ 4. Refresh Columns & Formulas", "E2E8F0", "1E293B")
    ]

    for row_idx, label, text_color, bg_color in actions:
        cell = ws[f"J{row_idx}"]
        cell.value = label
        cell.font = Font(name=FONT_NAME, size=10, bold=True, color=text_color, underline="single")
        cell.fill = PatternFill(start_color=bg_color, end_color=bg_color, fill_type="solid")
        cell.alignment = Alignment(horizontal="left", vertical="center", indent=1, wrap_text=True)
        cell.border = BORDER_CELL

    # J15: Export Directory Header
    ws["J15"] = "Export Directory (.xlsx sync):"
    ws["J15"].font = Font(name=FONT_NAME, size=10, bold=True, color="FFFFFF")
    ws["J15"].fill = FILL_HEADER_DARK
    ws["J15"].alignment = ALIGN_LEFT

    # J16: Export Directory Cell
    existing_j16 = str(ws["J16"].value or ws["I16"].value or ws["H16"].value or "").strip()
    if export_dir:
        exp_val = export_dir
    elif existing_j16 and existing_j16.lower() != "none" and not existing_j16.startswith("<"):
        exp_val = existing_j16
    else:
        local_cand = os.path.expanduser(r"~\OneDrive\Documents\Warframe")
        if os.path.exists(local_cand):
            exp_val = local_cand
        else:
            exp_val = DEFAULT_EXPORT_DIR_PLACEHOLDER

    ws["J16"] = exp_val
    ws["J16"].font = Font(name=FONT_NAME, size=9, color="F8FAFC")
    ws["J16"].fill = FILL_CARD_DARK
    ws["J16"].alignment = ALIGN_LEFT
    ws["J16"].border = BORDER_CELL

    # J17: Note
    ws["J17"] = "^ Clean .xlsx copy without macros is automatically exported here on every run."
    ws["J17"].font = Font(name=FONT_NAME, size=8, italic=True, color="94A3B8")
    ws["J17"].fill = FILL_CANVAS

    # Dark background for Column I gutter and Column J spacer rows
    for r in range(1, max_r + 5):
        i_gutter = ws.cell(row=r, column=9)
        if r != 2 and not i_gutter.value:
            i_gutter.fill = FILL_CANVAS
        j_cell = ws.cell(row=r, column=10)
        if not j_cell.value:
            j_cell.fill = FILL_CANVAS


def adjust_column_widths(ws) -> None:
    """Sets optimal column widths for clean readability with Stock, Visible, and Sidebar."""
    for col, w in COLUMN_WIDTHS.items():
        ws.column_dimensions[col].width = w
