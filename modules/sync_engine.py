"""
Warframe Market Sell Sync - Sync & Rollover Engine
==================================================
Manages order synchronization from market into Excel, formula refreshes,
session rollover to all-time revenue, and stock-ascending sorting.
"""

import os
import openpyxl
from typing import Dict, List, Tuple, Optional

from .config import (
    SHEET_NAME, FONT_TOTAL, FONT_TOTAL_LABEL, FILL_TOTAL, BORDER_TOTAL
)
from .sheet_layout import (
    get_open_excel_workbook, migrate_sheet_layout_if_needed,
    initialize_sheet_structure, apply_row_formulas_and_styling,
    populate_sidebar_column, adjust_column_widths
)
from .vba_manager import inject_vba_and_shapes, export_clean_xlsx


def sync_market_orders_to_excel(
    excel_file: str,
    orders: List[dict],
    username: str = "",
    jwt_token: str = ""
) -> Tuple[int, int]:
    """Synchronizes active sell orders into Excel with formulas, stock, visibility, and stock-ascending sort."""
    app, wb_com = get_open_excel_workbook(excel_file)
    was_open = False
    if wb_com:
        was_open = True
        try:
            wb_com.Close(SaveChanges=True)
        except Exception:
            pass

    is_new = not os.path.exists(excel_file)
    if not is_new:
        try:
            wb = openpyxl.load_workbook(excel_file, keep_vba=True)
            ws = wb[SHEET_NAME] if SHEET_NAME in wb.sheetnames else wb.active
        except Exception:
            is_new = True
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = SHEET_NAME
    else:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = SHEET_NAME

    # Migrate older sheet layouts if needed
    migrate_sheet_layout_if_needed(ws)
    initialize_sheet_structure(ws)

    # Read existing items (Rows 3+)
    existing_items: Dict[str, dict] = {}
    for r in range(3, ws.max_row + 1):
        name_val = ws.cell(row=r, column=1).value
        if name_val is not None:
            clean = str(name_val).strip()
            if clean.lower() == "total":
                continue

            p_val = ws.cell(row=r, column=2).value
            try:
                price = int(p_val)
            except (ValueError, TypeError):
                price = 0

            # Base stock from Col C formula (=base-E{r})
            c_val = str(ws.cell(row=r, column=3).value or "").strip()
            base_stk = 1
            if c_val.startswith("="):
                parts = c_val[1:].split("-")
                if parts:
                    try:
                        base_stk = int(parts[0].strip())
                    except ValueError:
                        base_stk = 1
            else:
                try:
                    base_stk = int(float(c_val))
                except ValueError:
                    base_stk = 1

            # Visible from Col D
            d_val = ws.cell(row=r, column=4).value
            if d_val is None or str(d_val).strip() == "":
                is_vis = True
            else:
                d_str = str(d_val).strip()
                is_vis = d_str in ("☑", "True", "true", "TRUE", "1", "yes", "Yes")

            # Col E (Current Session) and Col F (All Time)
            try:
                curr_sold = int(ws.cell(row=r, column=5).value or 0)
            except (ValueError, TypeError):
                curr_sold = 0

            try:
                all_sold = int(ws.cell(row=r, column=6).value or 0)
            except (ValueError, TypeError):
                all_sold = 0

            # Base All-Time Revenue from Col H (=base+G{r})
            h_val = str(ws.cell(row=r, column=8).value or "").strip()
            base_rev = 0
            if h_val.startswith("="):
                parts = h_val[1:].split("+")
                if parts and len(parts) >= 2:
                    try:
                        base_rev = int(parts[0].strip())
                    except ValueError:
                        base_rev = 0
                else:
                    base_rev = all_sold * price
            else:
                try:
                    base_rev = int(float(h_val))
                except ValueError:
                    base_rev = all_sold * price

            existing_items[clean.lower()] = {
                "name": clean,
                "price": price,
                "base_stock": max(1, base_stk),
                "is_visible": is_vis,
                "curr_sold": max(0, curr_sold),
                "all_sold": max(0, all_sold),
                "base_all_time_rev": max(0, base_rev)
            }

    # Deduplicate incoming market orders by name (keep highest price if multiple)
    dedup: Dict[str, dict] = {}
    for o in orders:
        k = o["name"].strip().lower()
        if k not in dedup or o["price"] > dedup[k]["price"]:
            dedup[k] = o

    updated_cnt = 0
    added_cnt = 0

    for name_key, order in dedup.items():
        price = order["price"]
        order_qty = max(1, int(order.get("quantity", 1)))
        order_vis = bool(order.get("visible", True))
        display_name = order["name"].strip()

        if name_key in existing_items:
            old_price = existing_items[name_key]["price"]
            existing_items[name_key]["name"] = display_name
            existing_items[name_key]["price"] = price
            # If market order is hidden / delisted and Excel had 0 stock, preserve 0 stock
            if not order_vis and existing_items[name_key].get("base_stock", 1) <= 0:
                existing_items[name_key]["base_stock"] = 0
            else:
                existing_items[name_key]["base_stock"] = order_qty
            existing_items[name_key]["is_visible"] = order_vis
            if old_price != price:
                updated_cnt += 1
        else:
            existing_items[name_key] = {
                "name": display_name,
                "price": price,
                "base_stock": order_qty,
                "is_visible": order_vis,
                "curr_sold": 0,
                "all_sold": 0,
                "base_all_time_rev": 0
            }
            added_cnt += 1

    # Sort items by Stock ASCENDING, then Price DESCENDING when stock is same
    item_list = list(existing_items.values())
    item_list.sort(key=lambda x: (
        max(0, x["base_stock"] - x["curr_sold"]), # Stock ascending
        -x["price"],                              # Price descending
        x["name"].lower()
    ))

    # Clear old table data area (Cols 1 to 8)
    max_clean_row = max(ws.max_row or 0, len(item_list) + 10)
    for r in range(3, max_clean_row + 1):
        for c in range(1, 9):
            ws.cell(row=r, column=c).value = None

    # Write sorted items into rows 3..
    for idx, it in enumerate(item_list):
        r = 3 + idx
        ws.cell(row=r, column=1, value=it["name"])
        ws.cell(row=r, column=2, value=it["price"])
        ws.cell(row=r, column=5, value=it["curr_sold"])
        ws.cell(row=r, column=6, value=it["all_sold"])
        apply_row_formulas_and_styling(
            ws,
            r,
            base_stock=it["base_stock"],
            is_visible=it["is_visible"],
            base_all_time_rev=it["base_all_time_rev"],
            is_zebra=(r % 2 == 0)
        )

    last_data_row = 2 + len(item_list)
    if last_data_row < 3:
        last_data_row = 3

    # Total Row
    tot_row = last_data_row + 1
    ws.row_dimensions[tot_row].height = 24
    ws.cell(row=tot_row, column=1, value="Total").font = FONT_TOTAL_LABEL
    ws.cell(row=tot_row, column=1).alignment = openpyxl.styles.Alignment(horizontal="left", vertical="center")
    ws.cell(row=tot_row, column=1).border = BORDER_TOTAL
    ws.cell(row=tot_row, column=1).fill = FILL_TOTAL

    for c in range(2, 7):
        cell = ws.cell(row=tot_row, column=c)
        cell.value = None
        cell.border = BORDER_TOTAL
        cell.fill = FILL_TOTAL

    tot_g = ws.cell(row=tot_row, column=7)
    tot_g.value = f"=SUM(G3:G{last_data_row})"
    tot_g.font = FONT_TOTAL
    tot_g.alignment = openpyxl.styles.Alignment(horizontal="right", vertical="center")
    tot_g.number_format = "#,##0"
    tot_g.border = BORDER_TOTAL
    tot_g.fill = FILL_TOTAL

    tot_h = ws.cell(row=tot_row, column=8)
    tot_h.value = f"=SUM(H3:H{last_data_row})"
    tot_h.font = FONT_TOTAL
    tot_h.alignment = openpyxl.styles.Alignment(horizontal="right", vertical="center")
    tot_h.number_format = "#,##0"
    tot_h.border = BORDER_TOTAL
    tot_h.fill = FILL_TOTAL

    populate_sidebar_column(ws, username=username, jwt_token=jwt_token)
    adjust_column_widths(ws)

    if is_new and excel_file.lower().endswith(".xlsm"):
        temp_xlsx = excel_file[:-5] + "_temp.xlsx"
        wb.save(temp_xlsx)
        wb.close()
        inject_vba_and_shapes(temp_xlsx, output_xlsm=excel_file)
        if os.path.exists(temp_xlsx):
            try:
                os.remove(temp_xlsx)
            except Exception:
                pass
    else:
        wb.save(excel_file)
        wb.close()
        inject_vba_and_shapes(excel_file)

    if was_open and app:
        try:
            app.Workbooks.Open(os.path.abspath(excel_file))
        except Exception:
            pass

    export_clean_xlsx(excel_file)
    return updated_cnt, added_cnt


def update_columns_and_formulas(excel_file: str) -> int:
    """Refreshes sheet structure, dynamic formulas, column widths, sorts stock ascending/price descending, and updates VBA."""
    if not os.path.exists(excel_file):
        print(f"[!] File '{excel_file}' not found.")
        return 0

    app, wb_com = get_open_excel_workbook(excel_file)
    was_open = False
    if wb_com:
        was_open = True
        try:
            wb_com.Close(SaveChanges=True)
        except Exception:
            pass

    wb = openpyxl.load_workbook(excel_file, keep_vba=True)
    ws = wb[SHEET_NAME] if SHEET_NAME in wb.sheetnames else wb.active

    migrate_sheet_layout_if_needed(ws)
    initialize_sheet_structure(ws)

    # Read existing items
    items = []
    for r in range(3, ws.max_row + 1):
        name_val = ws.cell(row=r, column=1).value
        if not name_val or str(name_val).strip().lower() == "total":
            continue
        clean = str(name_val).strip()

        p_val = ws.cell(row=r, column=2).value
        try:
            price = int(p_val)
        except (ValueError, TypeError):
            price = 0

        c_val = str(ws.cell(row=r, column=3).value or "").strip()
        base_stk = 1
        if c_val.startswith("="):
            parts = c_val[1:].split("-")
            if parts:
                try:
                    base_stk = int(parts[0].strip())
                except ValueError:
                    base_stk = 1
        else:
            try:
                base_stk = int(float(c_val))
            except ValueError:
                base_stk = 1

        d_val = ws.cell(row=r, column=4).value
        if d_val is None or str(d_val).strip() == "":
            is_vis = True
        else:
            d_str = str(d_val).strip()
            is_vis = d_str in ("☑", "True", "true", "TRUE", "1", "yes", "Yes")

        try:
            curr_sold = int(ws.cell(row=r, column=5).value or 0)
        except (ValueError, TypeError):
            curr_sold = 0

        try:
            all_sold = int(ws.cell(row=r, column=6).value or 0)
        except (ValueError, TypeError):
            all_sold = 0

        # Base All-Time Revenue from Col H (=base+G{r})
        h_val = str(ws.cell(row=r, column=8).value or "").strip()
        base_rev = 0
        if h_val.startswith("="):
            parts = h_val[1:].split("+")
            if parts and len(parts) >= 2:
                try:
                    base_rev = int(parts[0].strip())
                except ValueError:
                    base_rev = 0
            else:
                base_rev = all_sold * price
        else:
            try:
                base_rev = int(float(h_val))
            except ValueError:
                base_rev = all_sold * price

        items.append({
            "name": clean,
            "price": price,
            "base_stock": max(0, base_stk),
            "is_visible": is_vis,
            "curr_sold": max(0, curr_sold),
            "all_sold": max(0, all_sold),
            "base_all_time_rev": max(0, base_rev)
        })

    # Sort items by Stock ASCENDING, then Price DESCENDING when stock is same
    items.sort(key=lambda x: (
        max(0, x["base_stock"] - x["curr_sold"]), # Stock ascending
        -x["price"],                              # Price descending
        x["name"].lower()
    ))

    # Clear old table area
    max_clean_row = max(ws.max_row or 0, len(items) + 10)
    for r in range(3, max_clean_row + 1):
        for c in range(1, 9):
            ws.cell(row=r, column=c).value = None

    # Write sorted items
    for idx, it in enumerate(items):
        r = 3 + idx
        ws.cell(row=r, column=1, value=it["name"])
        ws.cell(row=r, column=2, value=it["price"])
        ws.cell(row=r, column=5, value=it["curr_sold"])
        ws.cell(row=r, column=6, value=it["all_sold"])
        apply_row_formulas_and_styling(
            ws,
            r,
            base_stock=it["base_stock"],
            is_visible=it["is_visible"],
            base_all_time_rev=it["base_all_time_rev"],
            is_zebra=(r % 2 == 0)
        )

    last_data_row = 2 + len(items)
    if last_data_row < 3:
        last_data_row = 3

    # Total Row
    tot_row = last_data_row + 1
    ws.row_dimensions[tot_row].height = 24
    ws.cell(row=tot_row, column=1, value="Total").font = FONT_TOTAL_LABEL
    ws.cell(row=tot_row, column=1).alignment = openpyxl.styles.Alignment(horizontal="left", vertical="center")
    ws.cell(row=tot_row, column=1).border = BORDER_TOTAL
    ws.cell(row=tot_row, column=1).fill = FILL_TOTAL

    for c in range(2, 7):
        cell = ws.cell(row=tot_row, column=c)
        cell.value = None
        cell.border = BORDER_TOTAL
        cell.fill = FILL_TOTAL

    tot_g = ws.cell(row=tot_row, column=7)
    tot_g.value = f"=SUM(G3:G{last_data_row})"
    tot_g.font = FONT_TOTAL
    tot_g.alignment = openpyxl.styles.Alignment(horizontal="right", vertical="center")
    tot_g.number_format = "#,##0"
    tot_g.border = BORDER_TOTAL
    tot_g.fill = FILL_TOTAL

    tot_h = ws.cell(row=tot_row, column=8)
    tot_h.value = f"=SUM(H3:H{last_data_row})"
    tot_h.font = FONT_TOTAL
    tot_h.alignment = openpyxl.styles.Alignment(horizontal="right", vertical="center")
    tot_h.number_format = "#,##0"
    tot_h.border = BORDER_TOTAL
    tot_h.fill = FILL_TOTAL

    populate_sidebar_column(ws)
    adjust_column_widths(ws)

    wb.save(excel_file)
    wb.close()

    inject_vba_and_shapes(excel_file)

    if was_open and app:
        try:
            app.Workbooks.Open(os.path.abspath(excel_file))
        except Exception:
            pass

    export_clean_xlsx(excel_file)
    return len(items)


def commit_session_to_all_time(excel_file: str) -> List[Tuple[str, int, int]]:
    """
    Rolls over Current Session quantities (Col E) to All Time (Col F),
    adds Current Session Revenue to All Time Revenue (Col H),
    decrements stock, unchecks visibility if stock falls from 1 to 0,
    resets Col E to 0, and sorts table by stock ascending and price descending.
    """
    if not os.path.exists(excel_file):
        print(f"[!] File '{excel_file}' not found.")
        return []

    app, wb_com = get_open_excel_workbook(excel_file)
    was_open = False
    if wb_com:
        was_open = True
        try:
            wb_com.Close(SaveChanges=True)
        except Exception:
            pass

    wb = openpyxl.load_workbook(excel_file, keep_vba=True)
    ws = wb[SHEET_NAME] if SHEET_NAME in wb.sheetnames else wb.active

    migrate_sheet_layout_if_needed(ws)

    items = []
    committed = []
    for r in range(3, ws.max_row + 1):
        name_val = ws.cell(row=r, column=1).value
        if not name_val or str(name_val).strip().lower() == "total":
            continue
        clean = str(name_val).strip()

        price_val = ws.cell(row=r, column=2).value or 0
        try:
            price = int(price_val)
        except (ValueError, TypeError):
            price = 0

        # Col 5 = Current Session, Col 6 = All Time
        try:
            curr_qty = int(ws.cell(row=r, column=5).value or 0)
        except (ValueError, TypeError):
            curr_qty = 0

        try:
            all_qty = int(ws.cell(row=r, column=6).value or 0)
        except (ValueError, TypeError):
            all_qty = 0

        # Base stock in Col 3 formula (=base-E{r})
        c_val = str(ws.cell(row=r, column=3).value or "").strip()
        base_stk = 1
        if c_val.startswith("="):
            parts = c_val[1:].split("-")
            if parts:
                try:
                    base_stk = int(parts[0].strip())
                except ValueError:
                    base_stk = 1
        else:
            try:
                base_stk = int(float(c_val))
            except ValueError:
                base_stk = 1

        # Visible in Col 4
        d_val = ws.cell(row=r, column=4).value
        if d_val is None or str(d_val).strip() == "":
            is_vis = True
        else:
            d_str = str(d_val).strip()
            is_vis = d_str in ("☑", "True", "true", "TRUE", "1", "yes", "Yes")

        # Base All-Time Revenue in Col 8 (=base+G{r})
        h_val = str(ws.cell(row=r, column=8).value or "").strip()
        base_all_time = 0
        if h_val.startswith("="):
            parts = h_val[1:].split("+")
            if parts and len(parts) >= 2:
                try:
                    base_all_time = int(parts[0].strip())
                except ValueError:
                    base_all_time = 0
            else:
                base_all_time = all_qty * price
        else:
            try:
                base_all_time = int(float(h_val))
            except ValueError:
                base_all_time = all_qty * price

        if curr_qty > 0:
            session_rev = curr_qty * price
            committed.append((clean, curr_qty, session_rev))
            all_qty += curr_qty
            base_all_time += session_rev  # Add Current Session Revenue to All Time Revenue!

            if base_stk - curr_qty <= 0:
                # When stock falls to 0: uncheck visibility checkbox and set base stock to 0
                is_vis = False
                new_base = 0
            else:
                new_base = base_stk - curr_qty
            curr_qty = 0
        else:
            new_base = base_stk

        items.append({
            "name": clean,
            "price": price,
            "base_stock": max(0, new_base),
            "is_visible": is_vis,
            "curr_sold": curr_qty,
            "all_sold": all_qty,
            "base_all_time_rev": max(0, base_all_time)
        })

    # Sort items by Stock ASCENDING, then Price DESCENDING when stock is same
    items.sort(key=lambda x: (
        max(0, x["base_stock"] - x["curr_sold"]), # Stock ascending
        -x["price"],                              # Price descending
        x["name"].lower()
    ))

    # Clear old table data area
    max_clean_row = max(ws.max_row or 0, len(items) + 10)
    for r in range(3, max_clean_row + 1):
        for c in range(1, 9):
            ws.cell(row=r, column=c).value = None

    # Re-write sorted rows
    for idx, it in enumerate(items):
        r = 3 + idx
        ws.cell(row=r, column=1, value=it["name"])
        ws.cell(row=r, column=2, value=it["price"])
        ws.cell(row=r, column=5, value=it["curr_sold"])
        ws.cell(row=r, column=6, value=it["all_sold"])
        apply_row_formulas_and_styling(
            ws,
            r,
            base_stock=it["base_stock"],
            is_visible=it["is_visible"],
            base_all_time_rev=it["base_all_time_rev"],
            is_zebra=(r % 2 == 0)
        )

    last_data_row = 2 + len(items)
    if last_data_row < 3:
        last_data_row = 3

    # Total Row
    tot_row = last_data_row + 1
    ws.row_dimensions[tot_row].height = 24
    ws.cell(row=tot_row, column=1, value="Total").font = FONT_TOTAL_LABEL
    ws.cell(row=tot_row, column=1).alignment = openpyxl.styles.Alignment(horizontal="left", vertical="center")
    ws.cell(row=tot_row, column=1).border = BORDER_TOTAL
    ws.cell(row=tot_row, column=1).fill = FILL_TOTAL

    for c in range(2, 7):
        cell = ws.cell(row=tot_row, column=c)
        cell.value = None
        cell.border = BORDER_TOTAL
        cell.fill = FILL_TOTAL

    tot_g = ws.cell(row=tot_row, column=7)
    tot_g.value = f"=SUM(G3:G{last_data_row})"
    tot_g.font = FONT_TOTAL
    tot_g.alignment = openpyxl.styles.Alignment(horizontal="right", vertical="center")
    tot_g.number_format = "#,##0"
    tot_g.border = BORDER_TOTAL
    tot_g.fill = FILL_TOTAL

    tot_h = ws.cell(row=tot_row, column=8)
    tot_h.value = f"=SUM(H3:H{last_data_row})"
    tot_h.font = FONT_TOTAL
    tot_h.alignment = openpyxl.styles.Alignment(horizontal="right", vertical="center")
    tot_h.number_format = "#,##0"
    tot_h.border = BORDER_TOTAL
    tot_h.fill = FILL_TOTAL

    populate_sidebar_column(ws)
    adjust_column_widths(ws)

    wb.save(excel_file)
    wb.close()

    inject_vba_and_shapes(excel_file)

    if was_open and app:
        try:
            app.Workbooks.Open(os.path.abspath(excel_file))
        except Exception:
            pass

    export_clean_xlsx(excel_file)
    return committed
