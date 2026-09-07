"""
Warframe Market Excel Automation Engine
=======================================
Single unified script managing Warframe.market synchronization, price updating,
session rollover, and Excel workbook automation.

Usage:
  python warframe_market.py                  # Interactive menu
  python warframe_market.py --sync           # Fetch sell orders and sync into Excel
  python warframe_market.py --push           # Read Excel prices and push changes to warframe.market
  python warframe_market.py --dry-run        # Preview price changes without pushing
  python warframe_market.py --commit         # Rollover session quantities (C -> D) and reset C to 0
  python warframe_market.py --update-columns # Refresh formatting, formulas, and VBA macro buttons
"""

import os
import sys
import json
import time
import base64
import shutil
import argparse
from typing import Dict, List, Optional, Tuple, Any

import requests
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

# Configure safe stdout encoding on Windows
if sys.platform == "win32" and hasattr(sys.stdout, "reconfigure"):
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except Exception:
        pass

# Constants
CACHE_FILE = "items_cache.json"
DEFAULT_EXCEL = "Warframe Sell Stats.xlsm"
SHEET_NAME = "Sheet1"
API_BASE_URL = "https://api.warframe.market/v2"
API_CALL_DELAY_SEC = 0.6
DEFAULT_EXPORT_DIR_PLACEHOLDER = "<PATH_TO_WARFRAME_FOLDER>"

BROWSER_HEADERS = {
    "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/148.0.0.0 Safari/537.36 Edg/148.0.0.0",
    "Accept": "application/json",
    "Content-Type": "application/json",
    "Platform": "pc",
    "Language": "en",
    "Crossplay": "true"
}

# Typography & Color Palette
FONT_NAME = "Segoe UI"
FONT_HEADER = Font(name=FONT_NAME, size=11, bold=True, color="FFFFFF")
FONT_SUBHEADER = Font(name=FONT_NAME, size=10, bold=True, color="FFFFFF")
FONT_DATA = Font(name=FONT_NAME, size=10, bold=False, color="1E293B")
FONT_TOTAL = Font(name=FONT_NAME, size=11, bold=True, color="0F172A")
FONT_MONO = Font(name="Consolas", size=9, color="0F172A")

FILL_HEADER_DARK = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
FILL_HEADER_BLUE = PatternFill(start_color="0369A1", end_color="0369A1", fill_type="solid")
FILL_HEADER_GREEN = PatternFill(start_color="047857", end_color="047857", fill_type="solid")
FILL_SUBHEADER = PatternFill(start_color="334155", end_color="334155", fill_type="solid")
FILL_ZEBRA_EVEN = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
FILL_TOTAL = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")
FILL_TOKEN = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")

ALIGN_LEFT = Alignment(horizontal="left", vertical="center")
ALIGN_CENTER = Alignment(horizontal="center", vertical="center")
ALIGN_RIGHT = Alignment(horizontal="right", vertical="center")

BORDER_THIN = Side(style="thin", color="CBD5E1")
BORDER_MEDIUM = Side(style="medium", color="475569")
BORDER_DOUBLE = Side(style="double", color="1E293B")
BORDER_AMBER = Side(style="medium", color="D97706")

BORDER_CELL = Border(left=BORDER_THIN, right=BORDER_THIN, top=BORDER_THIN, bottom=BORDER_THIN)
BORDER_TOTAL = Border(top=BORDER_MEDIUM, bottom=BORDER_DOUBLE, left=BORDER_THIN, right=BORDER_THIN)
BORDER_TOKEN_BOX = Border(left=BORDER_AMBER, right=BORDER_AMBER, top=BORDER_AMBER, bottom=BORDER_AMBER)


# ==============================================================================
# Credential & Excel Cell Configuration
# ==============================================================================

def get_open_excel_workbook(excel_file: str):
    """Returns (excel_app, workbook) if excel_file is currently open in Excel, else (None, None)."""
    try:
        import win32com.client
        excel = win32com.client.GetActiveObject("Excel.Application")
        target = os.path.abspath(excel_file).lower()
        for wb in excel.Workbooks:
            if os.path.abspath(wb.FullName).lower() == target:
                return excel, wb
    except Exception:
        pass
    return None, None


def get_credentials_from_excel(excel_file: str) -> Tuple[str, str]:
    """Reads Username from cell H2 and JWT Token from cell H4."""
    username = ""
    jwt_token = ""
    if not os.path.exists(excel_file):
        return username, jwt_token

    # Check active Excel first
    app, wb_com = get_open_excel_workbook(excel_file)
    if wb_com:
        try:
            ws = wb_com.Sheets(1)
            h2_val = str(ws.Range("H2").Value or "").strip()
            if h2_val and h2_val.lower() not in ("none", "enter your username here"):
                username = h2_val
            h4_val = str(ws.Range("H4").Value or "").strip()
            if h4_val and h4_val.lower() not in ("none", "paste your jwt token here"):
                if h4_val.lower().startswith("bearer "):
                    h4_val = h4_val[7:].strip()
                elif h4_val.lower().startswith("jwt "):
                    h4_val = h4_val[4:].strip()
                if len(h4_val) > 20:
                    jwt_token = h4_val
            return username, jwt_token
        except Exception:
            pass

    try:
        wb = openpyxl.load_workbook(excel_file, data_only=True, keep_vba=True)
        ws = wb[SHEET_NAME] if SHEET_NAME in wb.sheetnames else wb.active

        h2_val = str(ws["H2"].value or "").strip()
        if h2_val and h2_val.lower() not in ("none", "enter your username here"):
            username = h2_val

        h4_val = str(ws["H4"].value or "").strip()
        if h4_val and h4_val.lower() not in ("none", "paste your jwt token here"):
            if h4_val.lower().startswith("bearer "):
                h4_val = h4_val[7:].strip()
            elif h4_val.lower().startswith("jwt "):
                h4_val = h4_val[4:].strip()
            if len(h4_val) > 20:
                jwt_token = h4_val

        wb.close()
    except Exception:
        pass
    return username, jwt_token


def save_credentials_to_excel(excel_file: str, username: str = "", jwt_token: str = "") -> None:
    """Writes Username to cell H2 and JWT Token to cell H4 in the Excel workbook."""
    if not os.path.exists(excel_file):
        return

    # Check active Excel first to prevent file lock PermissionError
    app, wb_com = get_open_excel_workbook(excel_file)
    if wb_com:
        try:
            ws = wb_com.Sheets(1)
            if username:
                ws.Range("H2").Value = username
            if jwt_token:
                clean_tok = jwt_token.strip()
                if clean_tok.lower().startswith("bearer "):
                    clean_tok = clean_tok[7:].strip()
                elif clean_tok.lower().startswith("jwt "):
                    clean_tok = clean_tok[4:].strip()
                ws.Range("H4").Value = clean_tok
            wb_com.Save()
            return
        except Exception:
            pass

    try:
        wb = openpyxl.load_workbook(excel_file, keep_vba=True)
        ws = wb[SHEET_NAME] if SHEET_NAME in wb.sheetnames else wb.active

        if username:
            ws["H2"] = username
        if jwt_token:
            clean_tok = jwt_token.strip()
            if clean_tok.lower().startswith("bearer "):
                clean_tok = clean_tok[7:].strip()
            elif clean_tok.lower().startswith("jwt "):
                clean_tok = clean_tok[4:].strip()
            ws["H4"] = clean_tok

        wb.save(excel_file)
        wb.close()
    except Exception as e:
        print(f"[!] Warning: Could not update credentials in Excel: {e}")


def prompt_jwt_token(excel_file: str) -> str:
    """Prompts user for a fresh JWT token and writes it directly into Excel cell H4."""
    print("\n" + "=" * 70)
    print("  JWT TOKEN REQUIRED / EXPIRED - Warframe.market Authentication")
    print("=" * 70)
    print("How to get a fresh token from your browser (Chrome / Edge):")
    print("  1. Open https://warframe.market (ensure you are logged in)")
    print("  2. Press F12 -> Go to 'Application' or 'Storage' tab")
    print("  3. Under Storage -> Cookies -> https://warframe.market")
    print("  4. Find the cookie named 'JWT'")
    print("  5. Copy the value (starts with 'eyJ...')")
    print("  Note: You can also paste this directly into cell H4 in Excel!")
    print("=" * 70)

    token = input("\nPaste your JWT token here: ").strip()
    if token.lower().startswith("bearer "):
        token = token[7:].strip()
    elif token.lower().startswith("jwt "):
        token = token[4:].strip()

    if token:
        save_credentials_to_excel(excel_file, jwt_token=token)
        print(f"[+] Token saved to cell H4 of '{excel_file}'.")
    return token


def get_export_dir_from_excel(excel_file: str) -> Optional[str]:
    """Reads Export Directory from cell H16."""
    if not os.path.exists(excel_file):
        return None

    # Check active Excel first
    app, wb_com = get_open_excel_workbook(excel_file)
    if wb_com:
        try:
            ws = wb_com.Sheets(1)
            h16_val = str(ws.Range("H16").Value or "").strip()
            if h16_val and not h16_val.startswith("<") and h16_val.lower() not in ("none", ""):
                return h16_val
        except Exception:
            pass

    try:
        wb = openpyxl.load_workbook(excel_file, data_only=True, keep_vba=True)
        ws = wb[SHEET_NAME] if SHEET_NAME in wb.sheetnames else wb.active
        h16_val = str(ws["H16"].value or "").strip()
        wb.close()
        if h16_val and not h16_val.startswith("<") and h16_val.lower() not in ("none", ""):
            return h16_val
    except Exception:
        pass
    return None


def save_export_dir_to_excel(excel_file: str, export_dir: str) -> None:
    """Writes Export Directory path to cell H16."""
    if not os.path.exists(excel_file) or not export_dir:
        return

    app, wb_com = get_open_excel_workbook(excel_file)
    if wb_com:
        try:
            ws = wb_com.Sheets(1)
            ws.Range("H16").Value = export_dir.strip()
            wb_com.Save()
            return
        except Exception:
            pass

    try:
        wb = openpyxl.load_workbook(excel_file, keep_vba=True)
        ws = wb[SHEET_NAME] if SHEET_NAME in wb.sheetnames else wb.active
        ws["H16"] = export_dir.strip()
        wb.save(excel_file)
        wb.close()
    except Exception as e:
        print(f"[!] Warning: Could not update export directory in Excel: {e}")


def export_clean_xlsx(excel_file: str, export_dir: Optional[str] = None) -> Optional[str]:
    """
    Exports the workbook as a clean .xlsx copy without VBA macros or macro shapes into export_dir.
    Returns the path to the exported .xlsx file if successful, or None.
    """
    if not os.path.exists(excel_file):
        return None

    if not export_dir:
        export_dir = get_export_dir_from_excel(excel_file)

    if not export_dir or export_dir.startswith("<") or export_dir.lower() in ("none", ""):
        return None

    try:
        os.makedirs(export_dir, exist_ok=True)
        base_name = os.path.splitext(os.path.basename(excel_file))[0]
        dest_file = os.path.join(export_dir, f"{base_name}.xlsx")

        # Save active open workbook first if open in Excel COM
        app, wb_com = get_open_excel_workbook(excel_file)
        if wb_com:
            try:
                wb_com.Save()
            except Exception:
                pass

        # Load workbook with openpyxl (data_only=False to preserve formulas) without keep_vba
        wb = openpyxl.load_workbook(excel_file, data_only=False, keep_vba=False)
        for ws in wb.worksheets:
            if hasattr(ws, "_drawing") and ws._drawing:
                ws._drawing = None

        wb.save(dest_file)
        wb.close()
        print(f"[+] Clean .xlsx copy exported to: '{dest_file}'")
        return dest_file
    except Exception as e:
        print(f"[!] Note: Could not export clean .xlsx to '{export_dir}': {e}")
        return None


def get_auth_headers_and_cookies(jwt_token: str) -> Tuple[dict, dict]:
    """Builds Authorization header and cookies from JWT token."""
    headers = BROWSER_HEADERS.copy()
    headers["Authorization"] = f"Bearer {jwt_token}"

    csrf_token = None
    try:
        parts = jwt_token.split(".")
        if len(parts) >= 2:
            payload_b64 = parts[1] + "=" * ((4 - len(parts[1]) % 4) % 4)
            payload_data = json.loads(base64.urlsafe_b64decode(payload_b64).decode("utf-8"))
            csrf_token = payload_data.get("csrf_token")
    except Exception:
        pass

    if csrf_token:
        headers["x-csrftoken"] = csrf_token

    cookies = {"JWT": jwt_token}
    return headers, cookies


def is_jwt_expired(jwt_token: str) -> bool:
    """Checks whether the JWT token timestamp is expired."""
    if not jwt_token:
        return True
    try:
        parts = jwt_token.strip().split(".")
        if len(parts) >= 2:
            payload_b64 = parts[1] + "=" * ((4 - len(parts[1]) % 4) % 4)
            payload_data = json.loads(base64.urlsafe_b64decode(payload_b64).decode("utf-8"))
            exp = payload_data.get("exp")
            if exp and time.time() >= exp:
                return True
            return False
    except Exception:
        pass
    return True


def validate_jwt_token(jwt_token: str) -> bool:
    """Validates whether JWT token can successfully authenticate against the Warframe.market API."""
    if not jwt_token or is_jwt_expired(jwt_token):
        return False
    try:
        headers, cookies = get_auth_headers_and_cookies(jwt_token)
        resp = requests.get(f"{API_BASE_URL}/orders/my", headers=headers, cookies=cookies, timeout=8)
        return resp.status_code in (200, 201)
    except Exception:
        return False



# ==============================================================================
# Items Catalog Cache
# ==============================================================================

def load_items_catalog() -> Dict[str, dict]:
    """Loads cached items catalog or fetches fresh if missing."""
    if os.path.exists(CACHE_FILE):
        try:
            with open(CACHE_FILE, "r", encoding="utf-8") as f:
                cat = json.load(f)
                if cat and isinstance(cat, dict):
                    return cat
        except Exception:
            pass

    print("[*] Downloading items catalog from warframe.market...")
    try:
        resp = requests.get(f"{API_BASE_URL}/items", headers=BROWSER_HEADERS, timeout=15)
        resp.raise_for_status()
        items_list = resp.json().get("data", [])

        catalog = {}
        for it in items_list:
            i_id = it.get("id")
            en_name = it.get("i18n", {}).get("en", {}).get("name")
            if i_id and en_name:
                catalog[i_id] = {
                    "name": en_name,
                    "slug": it.get("slug", ""),
                    "tags": it.get("tags", []),
                    "maxRank": it.get("maxRank")
                }

        with open(CACHE_FILE, "w", encoding="utf-8") as f:
            json.dump(catalog, f)
        print(f"[+] Cached {len(catalog)} items.")
        return catalog
    except Exception as e:
        print(f"[!] Warning: Could not download items catalog: {e}")
        return {}


# ==============================================================================
# Warframe.market API Integration
# ==============================================================================

def fetch_orders(username: str, jwt_token: str, excel_file: str) -> List[dict]:
    """
    Fetches sell orders from warframe.market.
    Uses authenticated endpoint `/v2/orders/my` if JWT is available,
    otherwise uses public endpoint `/v2/orders/user/{username}`.
    """
    catalog = load_items_catalog()

    if jwt_token:
        headers, cookies = get_auth_headers_and_cookies(jwt_token)
        try:
            resp = requests.get(f"{API_BASE_URL}/orders/my", headers=headers, cookies=cookies, timeout=12)
            if resp.status_code in (200, 201):
                raw_orders = resp.json().get("data", [])
                orders = []
                for o in raw_orders:
                    if o.get("order_type") == "sell" or o.get("type") == "sell":
                        i_id = o.get("itemId") or o.get("item_id")
                        cat_item = catalog.get(i_id, {})
                        base_name = cat_item.get("name") or o.get("item", {}).get("slug", "Unknown Item")
                        rank = o.get("rank")
                        max_rank = cat_item.get("maxRank")
                        if rank is not None and max_rank and max_rank > 0:
                            display_name = f"{base_name} (Rank {rank})"
                        else:
                            display_name = base_name
                        orders.append({
                            "order_id": o.get("id"),
                            "item_id": i_id,
                            "name": display_name,
                            "base_name": base_name,
                            "price": int(o.get("platinum", 0)),
                            "quantity": int(o.get("quantity", 1)),
                            "rank": rank,
                            "visible": o.get("visible", True)
                        })
                return orders
            elif resp.status_code == 401:
                print("[!] Current JWT token has expired or is invalid.")
                fresh_token = prompt_jwt_token(excel_file)
                if fresh_token:
                    return fetch_orders(username, fresh_token, excel_file)
        except Exception as e:
            print(f"[!] Error with authenticated fetch: {e}")

    # Fallback to public orders by username
    if username:
        try:
            print(f"[*] Fetching public sell orders for user '{username}'...")
            resp = requests.get(f"{API_BASE_URL}/orders/user/{username}", headers=BROWSER_HEADERS, timeout=12)
            if resp.status_code == 200:
                raw_orders = resp.json().get("data", [])
                orders = []
                for o in raw_orders:
                    if o.get("order_type") == "sell" or o.get("type") == "sell":
                        i_id = o.get("itemId") or o.get("item_id")
                        cat_item = catalog.get(i_id, {})
                        base_name = cat_item.get("name") or o.get("item", {}).get("slug", "Unknown Item")
                        rank = o.get("rank")
                        max_rank = cat_item.get("maxRank")
                        if rank is not None and max_rank and max_rank > 0:
                            display_name = f"{base_name} (Rank {rank})"
                        else:
                            display_name = base_name
                        orders.append({
                            "order_id": o.get("id"),
                            "item_id": i_id,
                            "name": display_name,
                            "base_name": base_name,
                            "price": int(o.get("platinum", 0)),
                            "quantity": int(o.get("quantity", 1)),
                            "rank": rank,
                            "visible": o.get("visible", True)
                        })
                return orders
        except Exception as e:
            print(f"[!] Error with public fetch: {e}")

    return []


# ==============================================================================
# Excel Sheet Layout, Formulas & Column H Configuration
# ==============================================================================

def safe_merge(ws, cell_range: str) -> None:
    """Safely merges cells if range is not already merged."""
    existing = [str(r) for r in ws.merged_cells.ranges]
    if cell_range not in existing:
        ws.merge_cells(cell_range)


def initialize_sheet_structure(ws) -> None:
    """Builds standard 2-row table headers, merged groups, and labels."""
    ws["A1"] = "Item/Mod Name"
    ws["B1"] = "Price"
    ws["C1"] = "Quantity"
    ws["E1"] = "Revenue Generated"

    ws["C2"] = "Current Session"
    ws["D2"] = "All Time"
    ws["E2"] = "Current Session"
    ws["F2"] = "All Time"

    safe_merge(ws, "A1:A2")
    safe_merge(ws, "B1:B2")
    safe_merge(ws, "C1:D1")
    safe_merge(ws, "E1:F1")

    # Header cell styling
    for col in range(1, 7):
        cell1 = ws.cell(row=1, column=col)
        cell1.alignment = ALIGN_CENTER
        cell1.border = BORDER_CELL
        if col in (1, 2):
            cell1.fill = FILL_HEADER_DARK
            cell1.font = FONT_HEADER
        elif col in (3, 4):
            cell1.fill = FILL_HEADER_BLUE
            cell1.font = FONT_HEADER
        elif col in (5, 6):
            cell1.fill = FILL_HEADER_GREEN
            cell1.font = FONT_HEADER

        cell2 = ws.cell(row=2, column=col)
        cell2.alignment = ALIGN_CENTER
        cell2.border = BORDER_CELL
        if col in (1, 2):
            cell2.fill = FILL_HEADER_DARK
        else:
            cell2.fill = FILL_SUBHEADER
            cell2.font = FONT_SUBHEADER

    # G2: Linked Actions Header
    ws["G2"] = "Actions / Update All Time Revenue:"
    ws["G2"].font = Font(name=FONT_NAME, size=10, bold=True, color="0369A1", underline="single")


def apply_row_formulas_and_styling(ws, row_idx: int, is_zebra: bool = False) -> None:
    """Applies dynamic Excel formulas (=B*C and =B*D) and clean formatting."""
    ws[f"E{row_idx}"] = f"=B{row_idx}*C{row_idx}"
    ws[f"F{row_idx}"] = f"=B{row_idx}*D{row_idx}"

    row_fill = FILL_ZEBRA_EVEN if is_zebra else None
    for col_idx in range(1, 7):
        cell = ws.cell(row=row_idx, column=col_idx)
        cell.font = FONT_DATA
        cell.border = BORDER_CELL
        if row_fill:
            cell.fill = row_fill

        if col_idx == 1:
            cell.alignment = ALIGN_LEFT
        else:
            cell.alignment = ALIGN_RIGHT
            cell.number_format = "#,##0"


def populate_column_h(ws, username: str = "", jwt_token: str = "", export_dir: str = "") -> None:
    """
    Sets up Column H with stored credentials, export directory, and interactive action cells:
      H1: Header 'Warframe.market Account:'
      H2: Stored Username
      H3: Header 'Stored JWT Token (warframe.market):'
      H4: Stored JWT Token
      H5: Hint note
      H6: Header 'One-Click Macro Actions (Click Cell or Button to Run):'
      H7: '▶ 1. Sync from Market'
      H9: '⬆ 2. Push Prices to Market'
      H11: '🔄 3. Update All Time Revenue'
      H13: '⚡ 4. Refresh Columns & Formulas'
      H15: Header 'Export Directory (.xlsx sync):'
      H16: Target directory for automatic clean .xlsx export
      H17: Hint note
    """
    # H1: Username Header
    ws["H1"] = "Warframe.market Account:"
    ws["H1"].font = Font(name=FONT_NAME, size=10, bold=True, color="FFFFFF")
    ws["H1"].fill = FILL_HEADER_DARK
    ws["H1"].alignment = ALIGN_LEFT

    # H2: Username Cell
    existing_h2 = str(ws["H2"].value or "").strip()
    u_val = username if username else (existing_h2 if existing_h2 and existing_h2.lower() != "none" else "darksoulhunter2001")
    ws["H2"] = u_val
    ws["H2"].font = Font(name=FONT_NAME, size=10, bold=True, color="0F172A")
    ws["H2"].alignment = ALIGN_LEFT
    ws["H2"].fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    ws["H2"].border = BORDER_CELL

    # H3: JWT Header
    ws["H3"] = "Stored JWT Token (warframe.market):"
    ws["H3"].font = Font(name=FONT_NAME, size=10, bold=True, color="FFFFFF")
    ws["H3"].fill = FILL_HEADER_DARK
    ws["H3"].alignment = ALIGN_LEFT

    # H4: JWT Cell
    existing_h4 = str(ws["H4"].value or "").strip()
    tok_val = jwt_token if jwt_token else (existing_h4 if existing_h4 and len(existing_h4) > 20 else "Paste your JWT token here")
    ws["H4"] = tok_val
    ws["H4"].font = FONT_MONO
    ws["H4"].fill = FILL_TOKEN
    ws["H4"].alignment = ALIGN_LEFT
    ws["H4"].border = BORDER_TOKEN_BOX

    # H5: Note
    ws["H5"] = "^ Stored credentials. Read directly by Python. Paste fresh JWT above if expired."
    ws["H5"].font = Font(name=FONT_NAME, size=8, italic=True, color="64748B")

    # H6: Actions Header
    ws["H6"] = "One-Click Macro Actions (Click Cell or Button to Run):"
    ws["H6"].font = Font(name=FONT_NAME, size=10, bold=True, color="FFFFFF")
    ws["H6"].fill = FILL_HEADER_BLUE
    ws["H6"].alignment = ALIGN_LEFT

    # Action Items (Rows 7, 9, 11, 13)
    actions = [
        (7, "▶ 1. Sync from Market", "0284C7", "E0F2FE"),
        (9, "⬆ 2. Push Prices to Market", "059669", "D1FAE5"),
        (11, "🔄 3. Update All Time Revenue", "D97706", "FEF3C7"),
        (13, "⚡ 4. Refresh Columns & Formulas", "475569", "F1F5F9")
    ]

    for row_idx, label, text_color, bg_color in actions:
        cell = ws[f"H{row_idx}"]
        cell.value = label
        cell.font = Font(name=FONT_NAME, size=10, bold=True, color=text_color, underline="single")
        cell.fill = PatternFill(start_color=bg_color, end_color=bg_color, fill_type="solid")
        cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        cell.border = BORDER_CELL

    # H15: Export Directory Header
    ws["H15"] = "Export Directory (.xlsx sync):"
    ws["H15"].font = Font(name=FONT_NAME, size=10, bold=True, color="FFFFFF")
    ws["H15"].fill = FILL_HEADER_DARK
    ws["H15"].alignment = ALIGN_LEFT

    # H16: Export Directory Cell
    existing_h16 = str(ws["H16"].value or "").strip()
    if export_dir:
        exp_val = export_dir
    elif existing_h16 and existing_h16.lower() != "none" and not existing_h16.startswith("<"):
        exp_val = existing_h16
    else:
        # Default to local OneDrive Warframe folder if present, else placeholder
        local_cand = os.path.expanduser(r"~\OneDrive\Documents\Warframe")
        if os.path.exists(local_cand):
            exp_val = local_cand
        else:
            exp_val = DEFAULT_EXPORT_DIR_PLACEHOLDER

    ws["H16"] = exp_val
    ws["H16"].font = Font(name=FONT_NAME, size=9, color="0F172A")
    ws["H16"].fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    ws["H16"].alignment = ALIGN_LEFT
    ws["H16"].border = BORDER_CELL

    # H17: Note
    ws["H17"] = "^ Clean .xlsx copy without macros is automatically exported here on every run."
    ws["H17"].font = Font(name=FONT_NAME, size=8, italic=True, color="64748B")


def adjust_column_widths(ws) -> None:
    """Sets optimal column widths for clean readability."""
    widths = {
        "A": 34,
        "B": 12,
        "C": 16,
        "D": 14,
        "E": 18,
        "F": 16,
        "G": 28,
        "H": 80
    }
    for col, w in widths.items():
        ws.column_dimensions[col].width = w


# ==============================================================================
# VBA Macro & Shape Button Injection via Excel COM Automation
# ==============================================================================

VBA_MODULE_CODE = '''Option Explicit

Public Sub ExportNormalXLSX()
    On Error Resume Next
    Dim ws As Worksheet
    Set ws = ThisWorkbook.Sheets(1)
    Dim exportDir As String
    exportDir = Trim(CStr(ws.Range("H16").Value))
    If exportDir = "" Or InStr(exportDir, "<") > 0 Or LCase(exportDir) = "none" Then Exit Sub
    
    Dim fso As Object
    Set fso = CreateObject("Scripting.FileSystemObject")
    If Not fso.FolderExists(exportDir) Then
        fso.CreateFolder(exportDir)
    End If
    
    Dim baseName As String
    baseName = fso.GetBaseName(ThisWorkbook.Name)
    Dim destPath As String
    destPath = exportDir
    If Right(destPath, 1) <> "\\" Then destPath = destPath & "\\"
    destPath = destPath & baseName & ".xlsx"
    
    Dim prevAlerts As Boolean
    prevAlerts = Application.DisplayAlerts
    Application.DisplayAlerts = False
    
    Dim newWb As Workbook
    ws.Copy
    Set newWb = ActiveWorkbook
    
    ' Remove macro shapes / buttons from clean exported copy
    Dim shp As Shape
    For Each shp In newWb.Sheets(1).Shapes
        shp.Delete
    Next shp
    
    newWb.SaveAs Filename:=destPath, FileFormat:=51 ' 51 = xlOpenXMLWorkbook (.xlsx)
    newWb.Close SaveChanges:=False
    
    Application.DisplayAlerts = prevAlerts
End Sub

Public Sub SyncFromMarket(Optional ByVal ExtraArgs As String = "")
    Dim pyCmd As String
    Dim wsh As Object
    Set wsh = CreateObject("WScript.Shell")
    
    ThisWorkbook.Save
    Call ExportNormalXLSX
    
    Dim pyScript As String
    pyScript = Chr(34) & ThisWorkbook.Path & "\\warframe_market.py" & Chr(34)
    Dim xlFile As String
    xlFile = Chr(34) & ThisWorkbook.FullName & Chr(34)
    
    If ExtraArgs <> "" Then
        pyCmd = "cmd.exe /c cd /d " & Chr(34) & ThisWorkbook.Path & Chr(34) & " && python " & pyScript & " --sync --file " & xlFile & " " & ExtraArgs
        Call wsh.Run(pyCmd, 1, True)
    Else
        pyCmd = "cmd.exe /c cd /d " & Chr(34) & ThisWorkbook.Path & Chr(34) & " && python " & pyScript & " --sync --file " & xlFile
        Call wsh.Run(pyCmd, 1, False)
    End If
End Sub

Public Sub PushPricesToMarket(Optional ByVal ExtraArgs As String = "")
    Dim pyCmd As String
    Dim wsh As Object
    Set wsh = CreateObject("WScript.Shell")
    
    ThisWorkbook.Save
    Call ExportNormalXLSX
    
    Dim pyScript As String
    pyScript = Chr(34) & ThisWorkbook.Path & "\\warframe_market.py" & Chr(34)
    Dim xlFile As String
    xlFile = Chr(34) & ThisWorkbook.FullName & Chr(34)
    
    If ExtraArgs <> "" Then
        pyCmd = "cmd.exe /c cd /d " & Chr(34) & ThisWorkbook.Path & Chr(34) & " && python " & pyScript & " --push -y --file " & xlFile & " " & ExtraArgs
        Call wsh.Run(pyCmd, 1, True)
    Else
        pyCmd = "cmd.exe /c cd /d " & Chr(34) & ThisWorkbook.Path & Chr(34) & " && python " & pyScript & " --push -y --file " & xlFile
        Call wsh.Run(pyCmd, 1, False)
    End If
End Sub

Public Sub UpdateAllTimeRevenue()
    Dim ws As Worksheet
    Set ws = ThisWorkbook.Sheets(1)
    Dim r As Long, lastRow As Long
    Dim cVal As Double, dVal As Double, priceVal As Double
    Dim totalPlat As Double
    totalPlat = 0
    
    lastRow = ws.Cells(ws.Rows.Count, 1).End(xlUp).Row
    For r = 3 To lastRow
        If Trim(LCase(ws.Cells(r, 1).Value)) = "total" Then Exit For
        If ws.Cells(r, 1).Value <> "" Then
            cVal = Val(ws.Cells(r, 3).Value)
            dVal = Val(ws.Cells(r, 4).Value)
            priceVal = Val(ws.Cells(r, 2).Value)
            If cVal > 0 Then
                totalPlat = totalPlat + (cVal * priceVal)
                ws.Cells(r, 4).Value = dVal + cVal
                ws.Cells(r, 3).Value = 0
            End If
        End If
    Next r
    
    ThisWorkbook.Save
    Call ExportNormalXLSX
    If totalPlat > 0 Then
        Application.StatusBar = "All Time Revenue updated (" & Format(totalPlat, "#,##0") & " Plat added). Clean .xlsx copy exported."
    End If
End Sub

Public Sub RefreshColumnsAndFormulas(Optional ByVal ExtraArgs As String = "")
    Dim pyCmd As String
    Dim wsh As Object
    Set wsh = CreateObject("WScript.Shell")
    
    ThisWorkbook.Save
    Call ExportNormalXLSX
    
    Dim pyScript As String
    pyScript = Chr(34) & ThisWorkbook.Path & "\\warframe_market.py" & Chr(34)
    Dim xlFile As String
    xlFile = Chr(34) & ThisWorkbook.FullName & Chr(34)
    
    If ExtraArgs <> "" Then
        pyCmd = "cmd.exe /c cd /d " & Chr(34) & ThisWorkbook.Path & Chr(34) & " && python " & pyScript & " --update-columns --file " & xlFile & " " & ExtraArgs
        Call wsh.Run(pyCmd, 1, True)
    Else
        pyCmd = "cmd.exe /c cd /d " & Chr(34) & ThisWorkbook.Path & Chr(34) & " && python " & pyScript & " --update-columns --file " & xlFile
        Call wsh.Run(pyCmd, 1, False)
    End If
End Sub
'''

SHEET_EVENT_CODE = """Private Sub Worksheet_SelectionChange(ByVal Target As Range)
    On Error Resume Next
    If Target.Cells.CountLarge > 1 Then Exit Sub
    
    ' Trigger actions on cell clicks
    If Target.Row = 2 And Target.Column = 7 Then ' G2: Actions / Update All Time Revenue
        Application.EnableEvents = False
        Call WarframeMarket.UpdateAllTimeRevenue
        Range("A1").Select
        Application.EnableEvents = True
    ElseIf Target.Column = 8 Then
        Select Case Target.Row
            Case 7
                Application.EnableEvents = False
                Call WarframeMarket.SyncFromMarket
                Range("A1").Select
                Application.EnableEvents = True
            Case 9
                Application.EnableEvents = False
                Call WarframeMarket.PushPricesToMarket
                Range("A1").Select
                Application.EnableEvents = True
            Case 11
                Application.EnableEvents = False
                Call WarframeMarket.UpdateAllTimeRevenue
                Range("A1").Select
                Application.EnableEvents = True
            Case 13
                Application.EnableEvents = False
                Call WarframeMarket.RefreshColumnsAndFormulas
                Range("A1").Select
                Application.EnableEvents = True
        End Select
    End If
End Sub
"""

def inject_vba_and_shapes(input_path: str, output_xlsm: Optional[str] = None) -> bool:
    """
    Uses win32com to inject VBA module, sheet event handler, and macro shapes.
    Saves workbook as xlOpenXMLWorkbookMacroEnabled (52) .xlsm.
    """
    abs_input = os.path.abspath(input_path)
    abs_output = os.path.abspath(output_xlsm) if output_xlsm else abs_input

    try:
        import win32com.client
    except ImportError:
        print("[!] pywin32 not installed. Skipping VBA injection.")
        return False

    excel = None
    try:
        excel = win32com.client.Dispatch("Excel.Application")
        excel.Visible = False
        excel.DisplayAlerts = False

        wb = excel.Workbooks.Open(abs_input)
        ws = wb.Sheets(1)

        # 1. Inject or update WarframeMarket standard module
        vb_proj = wb.VBProject
        mod_found = False
        for comp in vb_proj.VBComponents:
            if comp.Name == "WarframeMarket":
                mod_found = True
                comp.CodeModule.DeleteLines(1, comp.CodeModule.CountOfLines)
                comp.CodeModule.AddFromString(VBA_MODULE_CODE)
                break
        if not mod_found:
            new_mod = vb_proj.VBComponents.Add(1)  # vbext_ct_StdModule
            new_mod.Name = "WarframeMarket"
            new_mod.CodeModule.AddFromString(VBA_MODULE_CODE)

        # 2. Inject Worksheet_SelectionChange event into Sheet1
        sheet_comp = None
        sheet_code_name = ws.CodeName or ws.Name
        for comp in vb_proj.VBComponents:
            if comp.Type == 100 and (comp.Name == sheet_code_name or comp.Name == ws.Name):
                sheet_comp = comp
                break
        if sheet_comp:
            code_mod = sheet_comp.CodeModule
            if code_mod.CountOfLines > 0:
                code_mod.DeleteLines(1, code_mod.CountOfLines)
            code_mod.AddFromString(SHEET_EVENT_CODE)

        # 3. Add or update OnAction Shape Buttons
        buttons_info = [
            ("Btn_Sync", 7, "▶  Sync from Market", "WarframeMarket.SyncFromMarket", (2, 132, 199)),
            ("Btn_Push", 9, "⬆  Push Prices to Market", "WarframeMarket.PushPricesToMarket", (5, 150, 105)),
            ("Btn_End", 11, "🔄  Update All Time Revenue", "WarframeMarket.UpdateAllTimeRevenue", (217, 119, 6)),
            ("Btn_Refresh", 13, "⚡  Refresh Formulas", "WarframeMarket.RefreshColumnsAndFormulas", (71, 85, 105))
        ]

        # Remove existing buttons if already present
        for shape in list(ws.Shapes):
            if shape.Name.startswith("Btn_"):
                shape.Delete()

        for btn_name, row_idx, text, macro_name, (r, g, b) in buttons_info:
            target_cell = ws.Cells(row_idx, 8)
            left = target_cell.Left + 4
            top = target_cell.Top + 2
            width = min(260, target_cell.Width - 8)
            height = target_cell.Height - 4

            shp = ws.Shapes.AddShape(5, left, top, width, height)  # 5 = msoShapeRoundedRectangle
            shp.Name = btn_name
            shp.TextFrame2.TextRange.Characters.Text = text
            shp.TextFrame2.TextRange.Font.Name = FONT_NAME
            shp.TextFrame2.TextRange.Font.Size = 9.5
            shp.TextFrame2.TextRange.Font.Bold = True
            shp.TextFrame2.TextRange.Font.Fill.ForeColor.RGB = 16777215  # White
            shp.TextFrame2.VerticalAnchor = 3  # msoAnchorMiddle
            shp.TextFrame2.TextRange.ParagraphFormat.Alignment = 2  # msoAlignCenter
            shp.Fill.Solid()
            shp.Fill.ForeColor.RGB = r + (g * 256) + (b * 65536)
            shp.Line.Visible = False
            shp.OnAction = macro_name

        if abs_output == abs_input:
            wb.Save()
        else:
            wb.SaveAs(abs_output, 52)  # 52 = xlOpenXMLWorkbookMacroEnabled
        wb.Close(False)
        wb = None
        return True
    except Exception as e:
        print(f"[!] Note on VBA injection: {e}")
        return False
    finally:
        try:
            if 'wb' in locals() and wb:
                wb.Close(False)
        except Exception:
            pass
        try:
            if excel:
                excel.Quit()
        except Exception:
            pass


# ==============================================================================
# Core Operations: Sync Orders, Commit Session, Refresh Columns
# ==============================================================================

def sync_market_orders_to_excel(
    excel_file: str,
    orders: List[dict],
    username: str = "",
    jwt_token: str = ""
) -> Tuple[int, int]:
    """Synchronizes active sell orders into Excel with dynamic formulas and styling."""
    app, wb_com = get_open_excel_workbook(excel_file)
    was_open = False
    if wb_com:
        was_open = True
        try:
            wb_com.Close(SaveChanges=True)
        except Exception:
            pass

    is_new = not os.path.exists(excel_file)
    if not is_new:
        try:
            wb = openpyxl.load_workbook(excel_file, keep_vba=True)
            ws = wb[SHEET_NAME] if SHEET_NAME in wb.sheetnames else wb.active
        except Exception:
            is_new = True
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = SHEET_NAME
    else:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = SHEET_NAME

    initialize_sheet_structure(ws)

    # Read existing items
    existing_items: Dict[str, int] = {}
    total_row_idx = None
    for r in range(3, ws.max_row + 1):
        name_val = ws.cell(row=r, column=1).value
        if name_val is not None:
            clean = str(name_val).strip()
            if clean.lower() == "total":
                total_row_idx = r
                continue
            existing_items[clean.lower()] = r

    if total_row_idx:
        ws.delete_rows(total_row_idx, 1)

    updated_cnt = 0
    added_cnt = 0
    last_data_row = 2
    for r in range(3, ws.max_row + 1):
        if ws.cell(row=r, column=1).value is not None:
            last_data_row = r

    # Deduplicate orders by name (keep highest price if multiple)
    dedup: Dict[str, dict] = {}
    for o in orders:
        k = o["name"].strip().lower()
        if k not in dedup or o["price"] > dedup[k]["price"]:
            dedup[k] = o

    for name_key, order in dedup.items():
        price = order["price"]
        display_name = order["name"].strip()

        if name_key in existing_items:
            r_idx = existing_items[name_key]
            old_price = ws.cell(row=r_idx, column=2).value
            ws.cell(row=r_idx, column=2, value=price)
            apply_row_formulas_and_styling(ws, r_idx, is_zebra=(r_idx % 2 == 0))
            if old_price != price:
                updated_cnt += 1
        else:
            last_data_row += 1
            ws.cell(row=last_data_row, column=1, value=display_name)
            ws.cell(row=last_data_row, column=2, value=price)
            if ws.cell(row=last_data_row, column=3).value is None:
                ws.cell(row=last_data_row, column=3, value=0)
            if ws.cell(row=last_data_row, column=4).value is None:
                ws.cell(row=last_data_row, column=4, value=0)
            apply_row_formulas_and_styling(ws, last_data_row, is_zebra=(last_data_row % 2 == 0))
            added_cnt += 1

    # Apply Total Row
    tot_row = last_data_row + 1 if last_data_row >= 3 else 4
    ws.cell(row=tot_row, column=1, value="Total").font = FONT_TOTAL
    ws.cell(row=tot_row, column=1).alignment = ALIGN_LEFT
    ws.cell(row=tot_row, column=1).border = BORDER_TOTAL
    ws.cell(row=tot_row, column=1).fill = FILL_TOTAL

    for c in range(2, 5):
        cell = ws.cell(row=tot_row, column=c)
        cell.border = BORDER_TOTAL
        cell.fill = FILL_TOTAL

    tot_e = ws.cell(row=tot_row, column=5)
    tot_e.value = f"=SUM(E3:E{last_data_row})"
    tot_e.font = FONT_TOTAL
    tot_e.alignment = ALIGN_RIGHT
    tot_e.number_format = "#,##0"
    tot_e.border = BORDER_TOTAL
    tot_e.fill = FILL_TOTAL

    tot_f = ws.cell(row=tot_row, column=6)
    tot_f.value = f"=SUM(F3:F{last_data_row})"
    tot_f.font = FONT_TOTAL
    tot_f.alignment = ALIGN_RIGHT
    tot_f.number_format = "#,##0"
    tot_f.border = BORDER_TOTAL
    tot_f.fill = FILL_TOTAL

    populate_column_h(ws, username=username, jwt_token=jwt_token)
    adjust_column_widths(ws)

    if is_new and excel_file.lower().endswith(".xlsm"):
        temp_xlsx = excel_file[:-5] + "_temp.xlsx"
        wb.save(temp_xlsx)
        wb.close()
        inject_vba_and_shapes(temp_xlsx, output_xlsm=excel_file)
        if os.path.exists(temp_xlsx):
            try:
                os.remove(temp_xlsx)
            except Exception:
                pass
    else:
        wb.save(excel_file)
        wb.close()
        inject_vba_and_shapes(excel_file)

    if was_open and app:
        try:
            app.Workbooks.Open(os.path.abspath(excel_file))
        except Exception:
            pass

    export_clean_xlsx(excel_file)
    return updated_cnt, added_cnt


def update_columns_and_formulas(excel_file: str) -> int:
    """Refreshes sheet structure, dynamic formulas, column widths, and VBA macros."""
    if not os.path.exists(excel_file):
        print(f"[!] File '{excel_file}' not found.")
        return 0

    app, wb_com = get_open_excel_workbook(excel_file)
    was_open = False
    if wb_com:
        was_open = True
        try:
            wb_com.Close(SaveChanges=True)
        except Exception:
            pass

    wb = openpyxl.load_workbook(excel_file, keep_vba=True)
    ws = wb[SHEET_NAME] if SHEET_NAME in wb.sheetnames else wb.active

    initialize_sheet_structure(ws)

    total_row_idx = None
    for r in range(3, ws.max_row + 1):
        val = ws.cell(row=r, column=1).value
        if val is not None and str(val).strip().lower() == "total":
            total_row_idx = r
            break

    if total_row_idx:
        ws.delete_rows(total_row_idx, 1)

    last_data_row = 2
    for r in range(3, ws.max_row + 1):
        if ws.cell(row=r, column=1).value is not None:
            last_data_row = r

    if last_data_row < 3:
        last_data_row = 3

    updated_rows = 0
    for r in range(3, last_data_row + 1):
        apply_row_formulas_and_styling(ws, r, is_zebra=(r % 2 == 0))
        updated_rows += 1

    tot_row = last_data_row + 1 if last_data_row >= 3 else 4
    ws.cell(row=tot_row, column=1, value="Total").font = FONT_TOTAL
    ws.cell(row=tot_row, column=1).alignment = ALIGN_LEFT
    ws.cell(row=tot_row, column=1).border = BORDER_TOTAL
    ws.cell(row=tot_row, column=1).fill = FILL_TOTAL

    for c in range(2, 5):
        cell = ws.cell(row=tot_row, column=c)
        cell.border = BORDER_TOTAL
        cell.fill = FILL_TOTAL

    tot_e = ws.cell(row=tot_row, column=5)
    tot_e.value = f"=SUM(E3:E{last_data_row})"
    tot_e.font = FONT_TOTAL
    tot_e.alignment = ALIGN_RIGHT
    tot_e.number_format = "#,##0"
    tot_e.border = BORDER_TOTAL
    tot_e.fill = FILL_TOTAL

    tot_f = ws.cell(row=tot_row, column=6)
    tot_f.value = f"=SUM(F3:F{last_data_row})"
    tot_f.font = FONT_TOTAL
    tot_f.alignment = ALIGN_RIGHT
    tot_f.number_format = "#,##0"
    tot_f.border = BORDER_TOTAL
    tot_f.fill = FILL_TOTAL

    populate_column_h(ws)
    adjust_column_widths(ws)

    wb.save(excel_file)
    wb.close()

    inject_vba_and_shapes(excel_file)

    if was_open and app:
        try:
            app.Workbooks.Open(os.path.abspath(excel_file))
        except Exception:
            pass

    export_clean_xlsx(excel_file)
    return updated_rows


def commit_session_to_all_time(excel_file: str) -> List[Tuple[str, int, int]]:
    """Rolls over Current Session quantities (Col C) to All Time (Col D) and resets Col C to 0."""
    if not os.path.exists(excel_file):
        print(f"[!] File '{excel_file}' not found.")
        return []

    app, wb_com = get_open_excel_workbook(excel_file)
    was_open = False
    if wb_com:
        was_open = True
        try:
            wb_com.Close(SaveChanges=True)
        except Exception:
            pass

    wb = openpyxl.load_workbook(excel_file, keep_vba=True)
    ws = wb[SHEET_NAME] if SHEET_NAME in wb.sheetnames else wb.active

    committed = []
    for r in range(3, ws.max_row + 1):
        name_val = ws.cell(row=r, column=1).value
        if not name_val or str(name_val).strip().lower() == "total":
            continue

        c_val = ws.cell(row=r, column=3).value or 0
        d_val = ws.cell(row=r, column=4).value or 0
        price_val = ws.cell(row=r, column=2).value or 0

        try:
            curr_qty = int(c_val)
        except (ValueError, TypeError):
            curr_qty = 0

        try:
            all_qty = int(d_val)
        except (ValueError, TypeError):
            all_qty = 0

        try:
            price = int(price_val)
        except (ValueError, TypeError):
            price = 0

        if curr_qty > 0:
            rev = curr_qty * price
            committed.append((str(name_val).strip(), curr_qty, rev))
            ws.cell(row=r, column=4, value=all_qty + curr_qty)
            ws.cell(row=r, column=3, value=0)

    wb.save(excel_file)
    wb.close()

    if was_open and app:
        try:
            app.Workbooks.Open(os.path.abspath(excel_file))
        except Exception:
            pass

    export_clean_xlsx(excel_file)
    return committed


# ==============================================================================
# Price Updater: Push Excel Prices to Warframe.market
# ==============================================================================

def read_excel_prices(excel_file: str) -> List[dict]:
    """Reads items and prices from Excel sheet."""
    if not os.path.exists(excel_file):
        return []

    app, wb_com = get_open_excel_workbook(excel_file)
    if wb_com:
        try:
            wb_com.Save()
        except Exception:
            pass

    wb = openpyxl.load_workbook(excel_file, data_only=True)
    ws = wb[SHEET_NAME] if SHEET_NAME in wb.sheetnames else wb.active

    items = []
    for r in range(3, ws.max_row + 1):
        name_val = ws.cell(row=r, column=1).value
        price_val = ws.cell(row=r, column=2).value
        if not name_val or str(name_val).strip().lower() == "total":
            continue
        try:
            p = int(price_val)
        except (ValueError, TypeError):
            continue

        items.append({
            "name": str(name_val).strip(),
            "price": p,
            "row": r
        })
    wb.close()
    return items


def diff_prices(excel_items: List[dict], market_orders: List[dict]) -> Tuple[List[dict], List[dict], List[dict]]:
    """Compares Excel prices with market orders."""
    market_lookup = {}
    for o in market_orders:
        market_lookup[o["name"].strip().lower()] = o
        base_k = o["base_name"].strip().lower()
        if base_k not in market_lookup:
            market_lookup[base_k] = o

    changed = []
    unchanged = []
    unmatched = []

    for ex in excel_items:
        ex_name = ex["name"].strip()
        ex_key = ex_name.lower()
        ex_price = ex["price"]

        order = market_lookup.get(ex_key)
        if not order:
            base_name = ex_name.split(" (Rank")[0].strip().lower()
            order = market_lookup.get(base_name)

        if not order:
            unmatched.append(ex)
            continue

        if order["price"] != ex_price:
            changed.append({
                "name": ex_name,
                "order_id": order["order_id"],
                "old_price": order["price"],
                "new_price": ex_price,
                "quantity": order["quantity"],
                "rank": order["rank"],
                "visible": order["visible"]
            })
        else:
            unchanged.append({
                "name": ex_name,
                "price": ex_price
            })

    return changed, unchanged, unmatched


def display_price_diff(changed: List[dict], unchanged: List[dict], unmatched: List[dict]) -> None:
    """Prints a formatted diff table of price changes."""
    print("\n" + "=" * 70)
    print("  WARFRAME MARKET PRICE UPDATE PREVIEW")
    print("=" * 70)

    if changed:
        print(f"\n  {len(changed)} PRICE CHANGE(S) DETECTED:")
        print("  " + "-" * 66)
        print(f"  {'Item Name':<36} {'Market':<10} {'-->':<5} {'Excel':<10}")
        print("  " + "-" * 66)
        for it in changed:
            diff = it["new_price"] - it["old_price"]
            sign = "+" if diff > 0 else "-"
            print(f"  {it['name']:<36} {it['old_price']:<10} {'-->':<5} {it['new_price']:<10} ({sign}{abs(diff)} plat)")
        print("  " + "-" * 66)
    else:
        print("\n  No price changes detected. All Excel prices match warframe.market!")

    if unchanged:
        print(f"  {len(unchanged)} item(s) already match market price.")
    if unmatched:
        print(f"  {len(unmatched)} item(s) in Excel were not found in market orders.")
    print("=" * 70)


def push_price_updates(
    jwt_token: str,
    changes: List[dict],
    excel_file: str
) -> Tuple[int, int]:
    """Pushes price updates via PATCH /v2/order/{order_id}."""
    headers, cookies = get_auth_headers_and_cookies(jwt_token)
    success = 0
    failed = 0

    for i, ch in enumerate(changes):
        order_id = ch["order_id"]
        new_price = ch["new_price"]
        item_name = ch["name"]

        payload = {
            "platinum": new_price,
            "quantity": ch["quantity"],
            "visible": ch["visible"]
        }
        if ch["rank"] is not None:
            payload["rank"] = ch["rank"]

        url = f"{API_BASE_URL}/order/{order_id}"
        try:
            resp = requests.patch(url, headers=headers, cookies=cookies, json=payload, timeout=10)
            if resp.status_code in (200, 201, 204):
                print(f"  [{i+1}/{len(changes)}] Updated '{item_name}': {ch['old_price']} -> {new_price} plat")
                success += 1
            elif resp.status_code == 401:
                print(f"  [!] JWT expired while updating '{item_name}'.")
                new_tok = prompt_jwt_token(excel_file)
                if new_tok:
                    jwt_token = new_tok
                    headers, cookies = get_auth_headers_and_cookies(jwt_token)
                    retry = requests.patch(url, headers=headers, cookies=cookies, json=payload, timeout=10)
                    if retry.status_code in (200, 201, 204):
                        print(f"  [{i+1}/{len(changes)}] Updated '{item_name}': {ch['old_price']} -> {new_price} plat")
                        success += 1
                        continue
                failed += 1
                break
            else:
                print(f"  [{i+1}/{len(changes)}] FAILED '{item_name}': HTTP {resp.status_code} - {resp.text[:100]}")
                failed += 1
        except Exception as e:
            print(f"  [{i+1}/{len(changes)}] FAILED '{item_name}': {e}")
            failed += 1

        if i < len(changes) - 1:
            time.sleep(API_CALL_DELAY_SEC)

    return success, failed


# ==============================================================================
# Workflows
# ==============================================================================

def run_sync_workflow(excel_file: str, username: str, jwt_token: str) -> None:
    """Sync orders from warframe.market into Excel."""
    print(f"\n[*] Connecting to Warframe.market...")
    orders = fetch_orders(username, jwt_token, excel_file)
    if not orders:
        print(f"[!] No active sell orders found for '{username}'.")
        return

    print(f"[+] Retrieved {len(orders)} active sell order(s).")
    print(f"[*] Updating spreadsheet '{excel_file}'...")
    updated_cnt, added_cnt = sync_market_orders_to_excel(excel_file, orders, username=username, jwt_token=jwt_token)
    print(f"[+] Success! Added {added_cnt} new item(s), refreshed {updated_cnt} price(s).")
    print(f"[+] Revenue formulas (=B*C and =B*D) and Totals row updated.")
    time.sleep(1.5)


def run_push_workflow(excel_file: str, username: str, jwt_token: str, dry_run: bool = False, auto_confirm: bool = False) -> None:
    """Read Excel prices, compare with market, and push updates."""
    if not jwt_token or not validate_jwt_token(jwt_token):
        print("\n[!] Current JWT token is missing, expired, or invalid.")
        jwt_token = prompt_jwt_token(excel_file)
        if not jwt_token or not validate_jwt_token(jwt_token):
            print("[!] Operation aborted. No valid JWT token provided.")
            return

    excel_items = read_excel_prices(excel_file)
    if not excel_items:
        print(f"[!] No items found in '{excel_file}'. Run --sync first.")
        return

    print(f"[*] Reading current orders from warframe.market...")
    market_orders = fetch_orders(username, jwt_token, excel_file)
    if not market_orders:
        print("[!] Could not fetch market orders.")
        return

    changed, unchanged, unmatched = diff_prices(excel_items, market_orders)
    display_price_diff(changed, unchanged, unmatched)

    if not changed:
        time.sleep(1.5)
        return

    if dry_run:
        print("\n[*] Dry-run mode enabled. No changes pushed.")
        return

    if not auto_confirm:
        ans = input(f"\nPush these {len(changed)} price change(s) to warframe.market? (y/N): ").strip().lower()
        if ans not in ("y", "yes"):
            print("[-] Cancelled by user.")
            return

    print(f"\n[*] Pushing {len(changed)} price update(s) to warframe.market...")
    success, failed = push_price_updates(jwt_token, changed, excel_file)
    print("\n" + "=" * 50)
    print(f"  PRICE UPDATE RESULT: {success} Succeeded, {failed} Failed")
    print("=" * 50)
    export_clean_xlsx(excel_file)
    time.sleep(1.5)


def run_commit_workflow(excel_file: str) -> None:
    """Update All Time Revenue: add Current Session quantities to All Time and reset to 0."""
    print(f"[*] Updating All Time Revenue in '{excel_file}'...")
    committed = commit_session_to_all_time(excel_file)
    if not committed:
        print("[-] Current Session had no quantities > 0. All Time Revenue unchanged.")
    else:
        print("\n" + "=" * 55)
        print("  ALL TIME REVENUE UPDATE SUMMARY")
        print("=" * 55)
        tot_plat = 0
        for name, qty, rev in committed:
            print(f"  - {name:<32} x{qty:<2} (+{rev} plat)")
            tot_plat += rev
        print("-" * 55)
        print(f"  Total Session Revenue: {tot_plat} Platinum")
        print(f"[+] Current Session quantities reset to 0.")
        print("=" * 55)


# ==============================================================================
# CLI Entry Point & Interactive Menu
# ==============================================================================

def main():
    parser = argparse.ArgumentParser(description="Warframe Market Excel Automation Engine")
    parser.add_argument("--sync", action="store_true", help="Fetch orders from warframe.market and sync to Excel")
    parser.add_argument("--push", action="store_true", help="Read Excel prices and push changes to warframe.market")
    parser.add_argument("--dry-run", action="store_true", help="Preview price changes without pushing")
    parser.add_argument("-y", "--yes", action="store_true", help="Skip confirmation prompt on price updates")
    parser.add_argument("--commit", action="store_true", help="Rollover Current Session quantities to All Time and reset to 0")
    parser.add_argument("--update-columns", action="store_true", help="Refresh formatting, formulas, and VBA macro buttons")
    parser.add_argument("--export-xlsx", action="store_true", help="Export clean .xlsx copy to configured directory")
    parser.add_argument("--user", type=str, help="Override Warframe.market Username")
    parser.add_argument("--token", type=str, help="Override / update JWT Token")
    parser.add_argument("--file", type=str, default=DEFAULT_EXCEL, help="Specify Excel file path")

    args = parser.parse_args()
    excel_file = args.file

    # Load credentials directly from Excel cells H2 & H4
    excel_user, excel_jwt = get_credentials_from_excel(excel_file)
    username = (args.user or excel_user or "darksoulhunter2001").strip()
    jwt_token = (args.token or excel_jwt or "").strip()

    if args.user or args.token:
        save_credentials_to_excel(excel_file, username=args.user, jwt_token=args.token)

    # Command Line Flags Execution
    if args.sync:
        run_sync_workflow(excel_file, username, jwt_token)
        return
    elif args.push:
        run_push_workflow(excel_file, username, jwt_token, dry_run=False, auto_confirm=args.yes)
        return
    elif args.dry_run:
        run_push_workflow(excel_file, username, jwt_token, dry_run=True, auto_confirm=True)
        return
    elif args.commit:
        run_commit_workflow(excel_file)
        return
    elif args.update_columns:
        cnt = update_columns_and_formulas(excel_file)
        print(f"[+] Refreshed {cnt} row(s) with formulas and formatting.")
        time.sleep(1.2)
        return
    elif args.export_xlsx:
        res = export_clean_xlsx(excel_file)
        if res:
            print(f"[+] Clean .xlsx exported: {res}")
        else:
            print(f"[!] Export failed or no valid export directory configured.")
        time.sleep(1.2)
        return

    # Interactive Menu Mode
    while True:
        excel_user, excel_jwt = get_credentials_from_excel(excel_file)
        u_display = excel_user or username or "[Not Set]"
        tok_display = "[Set in H4]" if (excel_jwt or jwt_token) else "[Not Set]"

        print("\n" + "=" * 60)
        print("       WARFRAME MARKET EXCEL AUTOMATION")
        print("=" * 60)
        print(f" Target File: {excel_file}")
        print(f" Account    : {u_display} | JWT: {tok_display}")
        print("-" * 60)
        print(" 1. Sync Market Sell Orders to Excel (Add missing items)")
        print(" 2. Push Excel Price Changes to Warframe.market")
        print(" 3. Preview Price Changes (Dry Run)")
        print(" 4. Update All Time Revenue (Commit C -> D & Reset to 0)")
        print(" 5. Refresh Columns, Formulas & Macro Buttons")
        print(" 6. Set / Update Browser JWT Token")
        print(" 7. Exit")
        print("=" * 60)

        choice = input("Select an option (1-7): ").strip()
        if choice == "1":
            run_sync_workflow(excel_file, username, jwt_token)
        elif choice == "2":
            run_push_workflow(excel_file, username, jwt_token)
        elif choice == "3":
            run_push_workflow(excel_file, username, jwt_token, dry_run=True)
        elif choice == "4":
            run_commit_workflow(excel_file)
        elif choice == "5":
            cnt = update_columns_and_formulas(excel_file)
            print(f"[+] Refreshed {cnt} row(s) with formulas and formatting.")
        elif choice == "6":
            new_tok = prompt_jwt_token(excel_file)
            if new_tok:
                jwt_token = new_tok
        elif choice == "7":
            print("Exiting. Happy trading, Tenno!")
            break
        else:
            print("[!] Invalid choice. Please enter 1-7.")


if __name__ == "__main__":
    main()
